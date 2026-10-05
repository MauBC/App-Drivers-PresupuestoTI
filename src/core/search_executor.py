from pathlib import Path
from datetime import datetime
import unicodedata
import re
from copy import deepcopy
from uuid import uuid4

import pandas as pd
from openpyxl import load_workbook
from rapidfuzz import process, fuzz
from src.core.excel_safety import (
    captured_sources, attach_capture, validate_configured_headers,
    atomic_write, atomic_update,
)


ESTADO_EXACTO = "EXACTO"
ESTADO_APROXIMADO = "APROXIMADO"
ESTADO_NO_ENCONTRADO = "NO ENCONTRADO"


def normalize_text(value) -> str:
    if value is None:
        return ""

    try:
        if pd.isna(value):
            return ""
    except Exception:
        pass

    text = str(value).strip()

    if not text:
        return ""

    if text.upper() in {"NAN", "NONE", "NULL", "NA", "N/A"}:
        return ""

    if text.endswith(".0"):
        possible_number = text[:-2]
        if possible_number.isdigit():
            text = possible_number

    text = unicodedata.normalize("NFKD", text)
    text = "".join(ch for ch in text if not unicodedata.combining(ch))
    text = text.upper()

    # Limpieza tipo script original
    text = re.sub(r"[,\.;:\-_()/]+", " ", text)
    text = re.sub(r"\s+", " ", text)

    return text.strip()


def generate_name_variants(text: str) -> list[str]:
    text = normalize_text(text)

    if not text:
        return []

    parts = text.split()
    variants = [text]

    if len(parts) >= 2:
        variants.append(" ".join(parts[::-1]))

    return list(dict.fromkeys(variants))


def find_best_similarity(query: str, choices: list[str]):
    variants = generate_name_variants(query)

    if not variants or not choices:
        return None

    best_match = None
    best_score = -1
    best_index = None

    for variant in variants:
        match = process.extractOne(
            variant,
            choices,
            scorer=fuzz.token_sort_ratio,
        )

        if match is None:
            continue

        matched_text, score, choice_index = match

        if score > best_score:
            best_match = matched_text
            best_score = score
            best_index = choice_index

    if best_match is None:
        return None

    return best_match, best_score, best_index


def estado_to_validado(estado: str) -> int:
    if estado in {ESTADO_EXACTO, ESTADO_APROXIMADO}:
        return 1
    return 0


def make_unique_column_name(name: str, used: set[str], suffix: str) -> str:
    base = str(name).strip() if str(name).strip() else "COLUMNA"

    if base not in used:
        used.add(base)
        return base

    candidate = f"{base}_{suffix}"
    counter = 2

    while candidate in used:
        candidate = f"{base}_{suffix}_{counter}"
        counter += 1

    used.add(candidate)
    return candidate


def get_columns_to_read(config: dict) -> list[dict]:
    columns = list(config["columnas"])
    search_col = config["columna_busqueda"]

    exists = any(col["indice_0"] == search_col["indice_0"] for col in columns)
    if not exists:
        columns.append(search_col)

    unique = {}
    for col in columns:
        unique[col["indice_0"]] = col

    return list(unique.values())


def read_configured_excel(config: dict, role: str) -> list[dict]:
    file_path = Path(config["archivo"]["ruta"])

    if not file_path.exists():
        raise FileNotFoundError(f"[{role}] No existe el archivo: {file_path}")

    sheet_name = config["hoja"]["nombre_detectado"]
    header_row = int(config["fila_header"])
    columns = get_columns_to_read(config)

    workbook = load_workbook(file_path, read_only=True, data_only=True)
    try:
        if sheet_name not in workbook.sheetnames:
            raise ValueError(f"[{role}] No existe la hoja: {sheet_name}")

        ws = workbook[sheet_name]

        max_row = ws.max_row or 0
        start_row = header_row + 1

        if start_row > max_row:
            return []

        records = []

        indices_1 = [col["indice_1"] for col in columns]
        min_col = min(indices_1)
        max_col = max(indices_1)

        index_to_col = {col["indice_1"]: col for col in columns}

        for excel_row_number, row in enumerate(
            ws.iter_rows(
                min_row=start_row,
                max_row=max_row,
                min_col=min_col,
                max_col=max_col,
                values_only=True,
            ),
            start=start_row,
        ):
            record = {
                "_excel_row_number": excel_row_number,
                "_role": role,
            }

            has_any_value = False

            for offset, value in enumerate(row):
                col_index_1 = min_col + offset
                col_config = index_to_col.get(col_index_1)

                if col_config is None:
                    continue

                alias = col_config["alias"]
                record[alias] = value

                if value is not None and str(value).strip() != "":
                    has_any_value = True

            if has_any_value:
                records.append(record)

        return records
    finally:
        workbook.close()


def get_search_value(record: dict, search_col: dict):
    alias = search_col["alias"]
    return record.get(alias)


def build_output_row(
    template_record: dict,
    base_record: dict | None,
    config: dict,
    estado: str,
    score,
    valor_buscado,
    valor_encontrado,
    observacion: str,
) -> dict:
    used_columns = set()
    row = {}

    template_columns = config["template"]["columnas"]
    base_columns = config["base"]["columnas"]

    for col in template_columns:
        if not col.get("mostrar_salida", True):
            continue

        alias = col["alias"]
        output_name = make_unique_column_name(alias, used_columns, "TEMPLATE")
        row[output_name] = template_record.get(alias)

    if base_record is not None:
        for col in base_columns:
            if not col.get("mostrar_salida", True):
                continue

            alias = col["alias"]
            output_name = make_unique_column_name(alias, used_columns, "BASE")
            row[output_name] = base_record.get(alias)
    else:
        for col in base_columns:
            if not col.get("mostrar_salida", True):
                continue

            alias = col["alias"]
            output_name = make_unique_column_name(alias, used_columns, "BASE")
            row[output_name] = ""

    if config["salida"].get("incluir_columnas_tecnicas", True):
        row["ESTADO_MATCH"] = estado
        row["SCORE_MATCH"] = score
        row["VALIDADO"] = estado_to_validado(estado)
        row["VALOR_BUSCADO"] = valor_buscado
        row["VALOR_ENCONTRADO"] = valor_encontrado
        row["FUENTE_ARCHIVO"] = config["base"]["archivo"]["nombre"] if base_record is not None else ""
        row["FUENTE_HOJA"] = config["base"]["hoja"]["nombre_detectado"] if base_record is not None else ""
        row["TIPO_BUSQUEDA"] = config["match"]["tipo_busqueda"]
        row["OBSERVACION"] = observacion

    return row


def run_exact_search(config: dict, template_records: list[dict], base_records: list[dict]) -> list[dict]:
    template_search_col = config["template"]["columna_busqueda"]
    base_search_col = config["base"]["columna_busqueda"]

    base_index = {}

    for base_record in base_records:
        raw_value = get_search_value(base_record, base_search_col)
        key = normalize_text(raw_value)

        if not key:
            continue

        if key not in base_index:
            base_index[key] = base_record

    results = []

    for template_record in template_records:
        valor_buscado = get_search_value(template_record, template_search_col)
        key = normalize_text(valor_buscado)

        if not key:
            results.append(
                build_output_row(
                    template_record=template_record,
                    base_record=None,
                    config=config,
                    estado=ESTADO_NO_ENCONTRADO,
                    score=0,
                    valor_buscado=valor_buscado,
                    valor_encontrado="",
                    observacion="Valor de busqueda vacio en template",
                )
            )
            continue

        base_record = base_index.get(key)

        if base_record is None:
            results.append(
                build_output_row(
                    template_record=template_record,
                    base_record=None,
                    config=config,
                    estado=ESTADO_NO_ENCONTRADO,
                    score=0,
                    valor_buscado=valor_buscado,
                    valor_encontrado="",
                    observacion="No hubo coincidencia exacta",
                )
            )
            continue

        valor_encontrado = get_search_value(base_record, base_search_col)

        results.append(
            build_output_row(
                template_record=template_record,
                base_record=base_record,
                config=config,
                estado=ESTADO_EXACTO,
                score=100,
                valor_buscado=valor_buscado,
                valor_encontrado=valor_encontrado,
                observacion="Coincidencia exacta",
            )
        )

    return results


def run_similarity_search(config: dict, template_records: list[dict], base_records: list[dict]) -> list[dict]:
    template_search_col = config["template"]["columna_busqueda"]
    base_search_col = config["base"]["columna_busqueda"]

    score_exacto = int(config["match"]["score_exacto"])
    score_aproximado = int(config["match"]["score_aproximado"])

    base_choices = []
    base_choice_to_record_index = []

    for index, base_record in enumerate(base_records):
        raw_value = get_search_value(base_record, base_search_col)
        normalized = normalize_text(raw_value)

        if not normalized:
            continue

        base_choices.append(normalized)
        base_choice_to_record_index.append(index)

    results = []
    match_cache = {}

    for template_record in template_records:
        valor_buscado = get_search_value(template_record, template_search_col)
        query = normalize_text(valor_buscado)

        if not query:
            results.append(
                build_output_row(
                    template_record=template_record,
                    base_record=None,
                    config=config,
                    estado=ESTADO_NO_ENCONTRADO,
                    score=0,
                    valor_buscado=valor_buscado,
                    valor_encontrado="",
                    observacion="Valor de busqueda vacio en template",
                )
            )
            continue

        if not base_choices:
            results.append(
                build_output_row(
                    template_record=template_record,
                    base_record=None,
                    config=config,
                    estado=ESTADO_NO_ENCONTRADO,
                    score=0,
                    valor_buscado=valor_buscado,
                    valor_encontrado="",
                    observacion="La base no tiene valores validos para buscar",
                )
            )
            continue

        if query not in match_cache:
            match_cache[query] = find_best_similarity(query, base_choices)
        match = match_cache[query]

        if match is None:
            results.append(
                build_output_row(
                    template_record=template_record,
                    base_record=None,
                    config=config,
                    estado=ESTADO_NO_ENCONTRADO,
                    score=0,
                    valor_buscado=valor_buscado,
                    valor_encontrado="",
                    observacion="No hubo coincidencia por similaridad",
                )
            )
            continue

        matched_text, score, choice_index = match
        score = round(float(score), 2)

        base_record_index = base_choice_to_record_index[choice_index]
        base_record = base_records[base_record_index]
        valor_encontrado = get_search_value(base_record, base_search_col)

        if score >= score_exacto:
            estado = ESTADO_EXACTO
            observacion = "Coincidencia por similaridad en rango EXACTO"
        elif score >= score_aproximado:
            estado = ESTADO_APROXIMADO
            observacion = "Coincidencia por similaridad en rango APROXIMADO"
        else:
            estado = ESTADO_NO_ENCONTRADO
            base_record = None
            valor_encontrado = ""
            observacion = "Score menor al rango APROXIMADO"

        results.append(
            build_output_row(
                template_record=template_record,
                base_record=base_record,
                config=config,
                estado=estado,
                score=score,
                valor_buscado=valor_buscado,
                valor_encontrado=valor_encontrado,
                observacion=observacion,
            )
        )

    return results


def run_search(config: dict, progress=None) -> tuple[pd.DataFrame, dict]:
    if progress:
        progress("Capturando template y maestro...")
    with captured_sources({role: config[role]["archivo"]["ruta"] for role in ("template", "base")}) as captures:
        local = deepcopy(config)
        if progress:
            progress("Validando encabezados...")
        for role in captures:
            local[role]["archivo"]["ruta"] = captures[role].snapshot
            validate_configured_headers(local[role])
        if progress:
            progress("Procesando busqueda...")
        df, summary = _run_search(local)
        attach_capture(df, captures, "template")
        return df, summary


def _run_search(config: dict) -> tuple[pd.DataFrame, dict]:
    template_records = read_configured_excel(config["template"], "template")
    base_records = read_configured_excel(config["base"], "base")

    if not template_records:
        raise ValueError("El template no tiene registros para procesar")

    if not base_records:
        raise ValueError("La base no tiene registros para buscar")

    search_type = config["match"]["tipo_busqueda"]

    if search_type == "exacta":
        result_rows = run_exact_search(config, template_records, base_records)
    elif search_type == "similaridad":
        result_rows = run_similarity_search(config, template_records, base_records)
    else:
        raise ValueError(f"Tipo de busqueda no soportado: {search_type}")

    df = pd.DataFrame(result_rows)

    summary = {
        "template_rows": len(template_records),
        "base_rows": len(base_records),
        "result_rows": len(df),
        "exactos": int((df["ESTADO_MATCH"] == ESTADO_EXACTO).sum()) if "ESTADO_MATCH" in df.columns else 0,
        "aproximados": int((df["ESTADO_MATCH"] == ESTADO_APROXIMADO).sum()) if "ESTADO_MATCH" in df.columns else 0,
        "no_encontrados": int((df["ESTADO_MATCH"] == ESTADO_NO_ENCONTRADO).sum()) if "ESTADO_MATCH" in df.columns else 0,
    }

    return df, summary


def sanitize_sheet_name(sheet_name: str) -> str:
    clean = str(sheet_name).strip()
    clean = re.sub(r"[\[\]\:\*\?\/\\]", "_", clean)
    clean = clean[:31]

    if not clean:
        clean = "RESULTADO_BUSQUEDA"

    return clean


def apply_xlsxwriter_format(writer, df: pd.DataFrame, sheet_name: str):
    workbook = writer.book
    worksheet = writer.sheets[sheet_name]

    header_format = workbook.add_format({
        "bold": True,
        "bg_color": "#5E6F32",
        "font_color": "#FFFFFF",
        "border": 1,
    })

    for col_num, value in enumerate(df.columns.values):
        worksheet.write(0, col_num, value, header_format)
        width = min(max(len(str(value)) + 4, 12), 45)
        worksheet.set_column(col_num, col_num, width)

    worksheet.freeze_panes(1, 0)

    if len(df.columns) > 0:
        worksheet.autofilter(0, 0, len(df), len(df.columns) - 1)

    if "ESTADO_MATCH" in df.columns and len(df) > 0:
        estado_col = list(df.columns).index("ESTADO_MATCH")

        fmt_exacto = workbook.add_format({"bg_color": "#D9EAD3", "font_color": "#274E13"})
        fmt_aprox = workbook.add_format({"bg_color": "#FFF2CC", "font_color": "#7F6000"})
        fmt_no = workbook.add_format({"bg_color": "#F4CCCC", "font_color": "#990000"})

        worksheet.conditional_format(1, estado_col, len(df), estado_col, {
            "type": "text",
            "criteria": "containing",
            "value": ESTADO_EXACTO,
            "format": fmt_exacto,
        })
        worksheet.conditional_format(1, estado_col, len(df), estado_col, {
            "type": "text",
            "criteria": "containing",
            "value": ESTADO_APROXIMADO,
            "format": fmt_aprox,
        })
        worksheet.conditional_format(1, estado_col, len(df), estado_col, {
            "type": "text",
            "criteria": "containing",
            "value": ESTADO_NO_ENCONTRADO,
            "format": fmt_no,
        })

    if "VALIDADO" in df.columns and len(df) > 0:
        validado_col = list(df.columns).index("VALIDADO")
        worksheet.data_validation(1, validado_col, len(df), validado_col, {
            "validate": "integer",
            "criteria": "between",
            "minimum": 0,
            "maximum": 1,
            "input_title": "VALIDADO",
            "input_message": "Usa 1 para validado o 0 para reprocesar.",
            "error_title": "Valor invalido",
            "error_message": "Solo se permite 1 o 0.",
        })


def export_result(df: pd.DataFrame, config: dict) -> Path:
    output_mode = config["salida"]["modo"]
    sheet_name = sanitize_sheet_name(config["salida"]["nombre_hoja_resultado"])

    if output_mode == "excel_nuevo":
        template_path = Path(config["template"]["archivo"]["ruta"])
        output_dir = template_path.parent

        timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        output_path = output_dir / f"{template_path.stem}_resultado_busqueda_{timestamp}_{uuid4().hex[:8]}.xlsx"

        def serialize(temporary):
            with pd.ExcelWriter(temporary, engine="xlsxwriter") as writer:
                df.to_excel(writer, index=False, sheet_name=sheet_name)
                apply_xlsxwriter_format(writer, df, sheet_name)

        return atomic_write(output_path, serialize)

    if output_mode == "nueva_hoja_template":
        template_path = Path(config["template"]["archivo"]["ruta"])

        def modify(workbook):
            if sheet_name in workbook.sheetnames:
                del workbook[sheet_name]
            ws = workbook.create_sheet(sheet_name)
            for col_idx, col_name in enumerate(df.columns, start=1):
                ws.cell(row=1, column=col_idx).value = col_name
            for row_idx, row_values in enumerate(df.itertuples(index=False, name=None), start=2):
                for col_idx, value in enumerate(row_values, start=1):
                    ws.cell(row=row_idx, column=col_idx).value = value
            ws.freeze_panes = "A2"
            ws.auto_filter.ref = ws.dimensions
        return atomic_update(df, template_path, modify)

    raise ValueError(f"Modo de salida no soportado: {output_mode}")
