"""Reglas comunes para nombres de columnas, sin quitar tildes ni símbolos."""
import unicodedata


def clean_header(value) -> str:
    if value is None:
        return ""
    text = str(value).translate({ord(c): None for c in "\u200b\ufeff\u2060"})
    return unicodedata.normalize("NFC", text).strip()


def header_key(value) -> str:
    return clean_header(value).casefold()


def prepare_headers(values) -> list[str]:
    headers = [clean_header(value) or f"COL_{index}" for index, value in enumerate(values, 1)]
    seen = {}
    for index, header in enumerate(headers, 1):
        key = header_key(header)
        if key in seen:
            raise ValueError(
                f"Encabezados ambiguos: '{header}' en columnas {seen[key]} y {index}. "
                "Asigne nombres distintos antes de continuar."
            )
        seen[key] = index
    return headers


def align_dataframe_headers(df, file_path, sheet_name, header_row):
    # Leer los nombres originales: pandas renombra duplicados automáticamente.
    from openpyxl import load_workbook
    workbook = load_workbook(file_path, read_only=True, data_only=True)
    try:
        values = next(workbook[sheet_name].iter_rows(
            min_row=header_row, max_row=header_row,
            max_col=len(df.columns), values_only=True,
        ), ()) if len(df.columns) else ()
        df.columns = prepare_headers(values)
    finally:
        workbook.close()
    return df
