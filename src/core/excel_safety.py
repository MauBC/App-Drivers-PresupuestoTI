"""Captura y persistencia Excel. No contiene reglas de matching.

Los hashes detectan cambios observados; no son una transaccion con Excel/OneDrive.
El guard coordina instancias de esta aplicacion en el mismo equipo y usuario.
"""
from contextlib import contextmanager
from dataclasses import dataclass, asdict
from datetime import datetime, timezone
from hashlib import sha256
import logging
import os
from pathlib import Path
import shutil
import tempfile
import time
from zipfile import ZipFile

from openpyxl import load_workbook


logger = logging.getLogger(__name__)
CAPTURE_ATTRIBUTE = "excel_capture"


def configure_logging(directory=None):
    """Diagnostico local, fuera del repositorio y de las carpetas de Excel."""
    from logging.handlers import RotatingFileHandler
    directory = Path(directory or Path(tempfile.gettempdir()) / "presupuesto_excel_logs")
    directory.mkdir(parents=True, exist_ok=True)
    path = directory / f"excel_{os.getpid()}.log"
    if not logger.handlers:
        handler = RotatingFileHandler(path, maxBytes=1_000_000, backupCount=2, encoding="utf-8")
        handler.setFormatter(logging.Formatter("%(asctime)s %(levelname)s %(message)s"))
        logger.addHandler(handler)
        logger.setLevel(logging.INFO)
    return path


class ExcelError(RuntimeError):
    pass


class ExcelBusyError(ExcelError):
    def __init__(self, path):
        self.path = Path(path)
        super().__init__(f"No se puede acceder a {self.path.name} porque esta ocupado o no tiene permisos. "
                         "Cierre el archivo en Microsoft Excel y vuelva a intentarlo.")


class ExcelChangedError(ExcelError):
    def __init__(self, path, expected, actual):
        self.path, self.expected, self.actual = Path(path), expected, actual
        super().__init__(f"{self.path.name} cambio mientras se procesaba la informacion. "
                         "No se sobrescribio el archivo para evitar perder cambios externos. Vuelva a ejecutar el proceso.")
        logger.warning("Cambio externo ruta=%s esperado=%s actual=%s", path, expected, actual)


class ExcelInvalidError(ExcelError):
    pass


class ExcelHeadersChangedError(ExcelError):
    pass


@dataclass(frozen=True)
class Snapshot:
    original: str
    snapshot: str
    sha256: str
    size: int
    mtime_ns: int
    captured_at: str
    attempt: int


def file_hash(path):
    with Path(path).open("rb") as stream:
        digest = sha256()
        for block in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(block)
        return digest.hexdigest()


def current_hash(path):
    try:
        return file_hash(path)
    except FileNotFoundError:
        return None


def transient(error):
    return isinstance(error, ExcelBusyError) or (
        isinstance(error, OSError) and (
            getattr(error, "winerror", None) in {5, 32, 33}
            or isinstance(error, PermissionError)
        )
    )


def retry(operation, path, attempts=3, delay=.15):
    for attempt in range(1, attempts + 1):
        try:
            return operation()
        except (OSError, ExcelBusyError) as error:
            if not transient(error):
                raise
            logger.warning("Archivo ocupado ruta=%s intento=%s/%s detalle=%r", path, attempt, attempts, error)
            if attempt == attempts:
                raise ExcelBusyError(path) from error
            time.sleep(delay * attempt)


def validate_source(path):
    path = Path(path)
    if path.name.startswith("~$"):
        raise ExcelInvalidError("Seleccione el Excel original, no el archivo temporal ~$ de Office.")
    if not path.is_file():
        raise FileNotFoundError(f"No existe el archivo Excel: {path}")
    if path.suffix.lower() not in {".xlsx", ".xlsm"}:
        raise ExcelInvalidError("Solo se soportan archivos .xlsx o .xlsm")
    return path.resolve()


def validate_package(path):
    try:
        with ZipFile(path) as archive:
            if archive.testzip() is not None:
                raise ExcelInvalidError(f"El Excel {Path(path).name} tiene datos corruptos.")
            if not {"[Content_Types].xml", "xl/workbook.xml"}.issubset(archive.namelist()):
                raise ExcelInvalidError(f"{Path(path).name} no contiene un libro Excel valido.")
        workbook = load_workbook(path, read_only=True, data_only=False)
        try:
            if not workbook.sheetnames:
                raise ExcelInvalidError("El libro no tiene hojas.")
        finally:
            workbook.close()
    except (OSError, ExcelInvalidError):
        raise
    except Exception as error:
        raise ExcelInvalidError(f"No se pudo validar el Excel {Path(path).name}: {error}") from error


@contextmanager
def captured_sources(sources, attempts=3, delay=.15):
    """Reintenta la captura COMPLETA si cualquiera de las fuentes cambia."""
    originals = {role: validate_source(path) for role, path in sources.items()}
    with tempfile.TemporaryDirectory(prefix="presupuesto_capture_") as directory:
        captures = None
        for attempt in range(1, attempts + 1):
            try:
                stats = {role: path.stat() for role, path in originals.items()}
                before = {role: file_hash(path) for role, path in originals.items()}
                copies = {}
                for index, (role, path) in enumerate(originals.items()):
                    target = Path(directory) / f"{index}{path.suffix.lower()}"
                    shutil.copyfile(path, target)
                    copies[role] = target
                copied = {role: file_hash(path) for role, path in copies.items()}
                after = {role: file_hash(path) for role, path in originals.items()}
                changed = [role for role in originals if before[role] != copied[role] or before[role] != after[role]]
                if changed:
                    role = changed[0]
                    observed = copied[role] if copied[role] != before[role] else after[role]
                    raise ExcelChangedError(originals[role], before[role], observed)
                captures = {}
                for role, path in originals.items():
                    validate_package(copies[role])
                    stat = stats[role]
                    captures[role] = Snapshot(str(path), str(copies[role]), copied[role],
                                              copies[role].stat().st_size, stat.st_mtime_ns,
                                              datetime.now(timezone.utc).isoformat(), attempt)
                    logger.info("Snapshot estable %s", asdict(captures[role]))
                break
            except (ExcelChangedError, OSError) as error:
                if not isinstance(error, ExcelChangedError) and not transient(error):
                    raise
                logger.warning("Captura descartada intento=%s/%s detalle=%r", attempt, attempts, error)
                if attempt == attempts:
                    if isinstance(error, ExcelChangedError):
                        raise
                    failed_path = getattr(error, "filename", None) or next(iter(originals.values()))
                    raise ExcelBusyError(failed_path) from error
                time.sleep(delay * attempt)
        yield captures


def attach_capture(df, captures, destination_role):
    df.attrs[CAPTURE_ATTRIBUTE] = {
        "destination": captures[destination_role].original,
        "expected": captures[destination_role].sha256,
        "sources": {role: asdict(snapshot) for role, snapshot in captures.items()},
    }


def expected_destination(df, destination):
    capture = df.attrs.get(CAPTURE_ATTRIBUTE)
    path = Path(destination).resolve()
    if not capture or path != Path(capture["destination"]):
        raise ExcelError("Ejecute el procesamiento antes de actualizar el Excel original.")
    for role, source in capture["sources"].items():
        if role == "base" and path == Path(source["original"]):
            raise ExcelError("El maestro es de solo lectura. Seleccione un template diferente o genere un archivo nuevo.")
    return capture["expected"]


def assert_unchanged(path, expected):
    actual = current_hash(path)
    if actual != expected:
        raise ExcelChangedError(path, expected, actual)


def check_writable(path):
    """Señal Office + prueba real de acceso Windows; nunca modifica el libro."""
    path = Path(path)
    if path.with_name("~$" + path.name).exists():
        raise ExcelBusyError(path)
    if path.exists() and os.name == "nt":
        import ctypes
        from ctypes import wintypes
        kernel = ctypes.WinDLL("kernel32", use_last_error=True)
        create = kernel.CreateFileW
        create.argtypes = [wintypes.LPCWSTR, wintypes.DWORD, wintypes.DWORD, ctypes.c_void_p,
                           wintypes.DWORD, wintypes.DWORD, wintypes.HANDLE]
        create.restype = wintypes.HANDLE
        close = kernel.CloseHandle
        close.argtypes = [wintypes.HANDLE]
        close.restype = wintypes.BOOL
        handle = create(str(path.resolve()), 0xC0000000, 0, None, 3, 0x80, None)
        if handle == ctypes.c_void_p(-1).value:
            raise ctypes.WinError(ctypes.get_last_error())
        close(handle)


@contextmanager
def destination_guard(path):
    """Lock de SO: se libera al morir el proceso. No borrar el archivo de lock."""
    path = Path(path).resolve()
    lock_dir = Path(tempfile.gettempdir()) / "presupuesto_excel_locks"
    lock_dir.mkdir(exist_ok=True)
    key = sha256(os.path.normcase(str(path)).encode("utf-8")).hexdigest()
    with (lock_dir / (key + ".lock")).open("a+b") as stream:
        stream.seek(0, 2)
        if stream.tell() == 0:
            stream.write(b"0")
            stream.flush()
        stream.seek(0)
        try:
            if os.name == "nt":
                import msvcrt
                msvcrt.locking(stream.fileno(), msvcrt.LK_NBLCK, 1)
            else:
                import fcntl
                fcntl.flock(stream, fcntl.LOCK_EX | fcntl.LOCK_NB)
        except OSError as error:
            raise ExcelBusyError(path) from error
        try:
            yield
        finally:
            stream.seek(0)
            if os.name == "nt":
                msvcrt.locking(stream.fileno(), msvcrt.LK_UNLCK, 1)
            else:
                fcntl.flock(stream, fcntl.LOCK_UN)


def atomic_write(destination, serialize, expected=None):
    """Serialize recibe un temporal hermano. expected=None exige destino inexistente.

    El callback debe cerrar writers/workbooks antes de regresar. El lock cubre la
    construccion y publicacion. Un fallo anterior al replace conserva el original.
    """
    destination = Path(destination)
    destination.parent.mkdir(parents=True, exist_ok=True)
    with destination_guard(destination):
        retry(lambda: check_writable(destination), destination)
        assert_unchanged(destination, expected)
        descriptor, name = tempfile.mkstemp(prefix=".presupuesto_", suffix=destination.suffix, dir=destination.parent)
        os.close(descriptor)
        temporary = Path(name)
        try:
            serialize(temporary)
            validate_package(temporary)
            generated_hash = file_hash(temporary)
            with temporary.open("r+b") as stream:
                stream.flush()
                os.fsync(stream.fileno())
            def publish():
                check_writable(destination)
                assert_unchanged(destination, expected)
                if expected is None:
                    # Reserva no destructiva: un tercero que crea el nombre no se pierde.
                    try:
                        if os.name == "nt":
                            os.rename(temporary, destination)
                        else:
                            os.link(temporary, destination)
                            temporary.unlink()
                    except FileExistsError as error:
                        raise ExcelChangedError(destination, None, current_hash(destination)) from error
                else:
                    os.replace(temporary, destination)
            retry(publish, destination)
            logger.info("Excel guardado ruta=%s sha256=%s", destination, generated_hash)
        except Exception:
            logger.exception("Fallo de persistencia destino=%s temporal=%s", destination, temporary)
            raise
        finally:
            if temporary.exists():
                try:
                    retry(temporary.unlink, temporary)
                except ExcelBusyError:
                    logger.exception("No se pudo limpiar temporal=%s", temporary)
    return destination


def atomic_update(df, destination, modify):
    expected = expected_destination(df, destination)
    def serialize(temporary):
        workbook = load_workbook(destination, keep_vba=Path(destination).suffix.lower() == ".xlsm")
        try:
            modify(workbook)
            workbook.save(temporary)
        finally:
            workbook.close()
            if workbook.vba_archive is not None:
                workbook.vba_archive.close()
    return atomic_write(destination, serialize, expected)


def validate_configured_headers(config):
    """Valida las posiciones utilizadas sin impedir cambios en filas de datos."""
    from src.core.excel_inspector import read_headers
    info = read_headers(config["archivo"]["ruta"], config["hoja"]["nombre_detectado"], int(config["fila_header"]))
    for column in config["columnas"] + [config["columna_busqueda"]]:
        index = column["indice_0"]
        expected = column.get("header_detectado")
        if expected is not None and (index >= len(info["headers"]) or info["headers"][index] != expected):
            raise ExcelHeadersChangedError("Los encabezados cambiaron desde la configuracion. Vuelva a leerlos y revise las columnas.")
