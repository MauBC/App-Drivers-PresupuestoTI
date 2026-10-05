from pathlib import Path
from openpyxl import load_workbook
from openpyxl.utils.cell import column_index_from_string, get_column_letter


VALID_EXTENSIONS = {".xlsx", ".xlsm"}


def validate_excel_path(file_path: str) -> Path:
    path = Path(file_path)

    if not path.exists():
        raise FileNotFoundError(f"No existe el archivo: {path}")

    if path.suffix.lower() not in VALID_EXTENSIONS:
        raise ValueError("Solo se soportan archivos .xlsx o .xlsm")

    return path


def parse_sheet_ref(sheet_ref: str, sheet_names: list[str]) -> tuple[int, str]:
    value = str(sheet_ref).strip()

    if not value:
        raise ValueError("La hoja no puede estar vacia")

    if value.isdigit():
        pos = int(value)

        if pos < 1 or pos > len(sheet_names):
            raise ValueError(f"La posicion de hoja debe estar entre 1 y {len(sheet_names)}")

        index = pos - 1
        return index, sheet_names[index]

    for index, name in enumerate(sheet_names):
        if value.lower() == name.lower():
            return index, name

    raise ValueError(f"No se encontro la hoja: {value}")


class ExcelWorkbookSession:
    """
    Mantiene un Excel abierto en modo lectura para no abrirlo una y otra vez.

    Uso:
    session = ExcelWorkbookSession("archivo.xlsx")
    session.read_headers("1", 6)
    session.close()
    """

    def __init__(self, file_path: str):
        self.path = validate_excel_path(file_path)
        self.workbook = load_workbook(
            self.path,
            read_only=True,
            data_only=True,
        )
        self.sheet_names = list(self.workbook.sheetnames)

    def close(self):
        if self.workbook is not None:
            self.workbook.close()
            self.workbook = None

    def _ensure_open(self):
        if self.workbook is None:
            raise RuntimeError("La sesion del Excel esta cerrada")

    def parse_sheet(self, sheet_ref: str) -> tuple[int, str]:
        self._ensure_open()
        return parse_sheet_ref(sheet_ref, self.sheet_names)

    def read_sheet_sample(self, sheet_ref: str, max_rows: int = 5, max_cols: int = 8) -> dict:
        self._ensure_open()

        sheet_index, sheet_name = self.parse_sheet(sheet_ref)
        worksheet = self.workbook[sheet_name]

        max_row = worksheet.max_row or 0
        max_column = worksheet.max_column or 0

        rows = []
        limit_rows = min(max_rows, max_row)
        limit_cols = min(max_cols, max_column)

        if limit_rows > 0 and limit_cols > 0:
            for row in worksheet.iter_rows(
                min_row=1,
                max_row=limit_rows,
                min_col=1,
                max_col=limit_cols,
                values_only=True,
            ):
                rows.append(["" if value is None else str(value) for value in row])

        return {
            "sheet_index": sheet_index,
            "sheet_name": sheet_name,
            "max_row": max_row,
            "max_column": max_column,
            "rows": rows,
        }

    def read_headers(self, sheet_ref: str, header_row: int) -> dict:
        self._ensure_open()

        sheet_index, sheet_name = self.parse_sheet(sheet_ref)

        if header_row < 1:
            raise ValueError("La fila del header debe ser mayor o igual a 1")

        worksheet = self.workbook[sheet_name]
        max_row = worksheet.max_row or 0

        if header_row > max_row:
            raise ValueError(f"La fila header {header_row} supera el maximo de filas")

        row_iterator = worksheet.iter_rows(
            min_row=header_row,
            max_row=header_row,
            values_only=True,
        )

        raw_values = list(next(row_iterator, []))

        last_real_index = -1
        for i, value in enumerate(raw_values):
            if value is not None and str(value).strip() != "":
                last_real_index = i

        if last_real_index == -1:
            raise ValueError("No se encontraron headers en esa fila")

        raw_values = raw_values[: last_real_index + 1]

        headers = []
        for i, value in enumerate(raw_values, start=1):
            if value is None or str(value).strip() == "":
                headers.append(f"COL_{i}")
            else:
                headers.append(str(value).strip())

        preview_rows = []
        start_row = header_row + 1
        end_row = min(header_row + 5, max_row)

        if end_row >= start_row:
            for row in worksheet.iter_rows(
                min_row=start_row,
                max_row=end_row,
                min_col=1,
                max_col=len(headers),
                values_only=True,
            ):
                preview_rows.append(["" if value is None else str(value) for value in row])

        return {
            "sheet_index": sheet_index,
            "sheet_name": sheet_name,
            "header_row": header_row,
            "headers": headers,
            "first_5_headers": headers[:5],
            "first_30_headers": headers[:30],
            "preview_rows": preview_rows,
        }


def get_sheet_names(file_path: str) -> list[str]:
    session = ExcelWorkbookSession(file_path)
    try:
        return session.sheet_names
    finally:
        session.close()


def read_sheet_sample(file_path: str, sheet_ref: str, max_rows: int = 5, max_cols: int = 8) -> dict:
    session = ExcelWorkbookSession(file_path)
    try:
        return session.read_sheet_sample(sheet_ref, max_rows=max_rows, max_cols=max_cols)
    finally:
        session.close()


def read_headers(file_path: str, sheet_ref: str, header_row: int) -> dict:
    session = ExcelWorkbookSession(file_path)
    try:
        return session.read_headers(sheet_ref, header_row)
    finally:
        session.close()


def column_ref_to_index(column_ref: str) -> int:
    value = str(column_ref).strip().upper()

    if not value:
        raise ValueError("La columna no puede estar vacia")

    if value.isdigit():
        index_1_based = int(value)
    else:
        index_1_based = column_index_from_string(value)

    if index_1_based < 1:
        raise ValueError("La columna debe ser mayor o igual a 1")

    return index_1_based - 1


def resolve_column(headers: list[str], column_ref: str) -> dict:
    index_0 = column_ref_to_index(column_ref)

    if index_0 >= len(headers):
        raise ValueError(f"La columna {column_ref} esta fuera del rango de headers detectados")

    return {
        "posicion_usuario": str(column_ref).strip(),
        "letra_excel": get_column_letter(index_0 + 1),
        "indice_0": index_0,
        "indice_1": index_0 + 1,
        "header_detectado": headers[index_0],
    }
