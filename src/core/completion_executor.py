from pathlib import Path
from datetime import datetime
import re

import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import PatternFill, Font
from openpyxl.utils.cell import column_index_from_string
from rapidfuzz import process, fuzz

from src.core.excel_inspector import parse_sheet_ref
from src.core.search_executor import (
    ESTADO_EXACTO,
    ESTADO_APROXIMADO,
    ESTADO_NO_ENCONTRADO,
    normalize_text,
    find_best_similarity,
    read_configured_excel,
    get_search_value,
    sanitize_sheet_name,
    apply_xlsxwriter_format,
)


TECHNICAL_COLUMNS = {
    "ESTADO_MATCH",
    "SCORE_MATCH",
    "VALIDADO",
    "VALOR_BUSCADO",
    "VALOR_ENCONTRADO",
    "FUENTE_ARCHIVO",
    "FUENTE_HOJA",
    "TIPO_BUSQUEDA",
    "OBSERVACION",
}

REQUIRED_TECHNICAL_COLUMNS = {
    "ESTADO_MATCH",
    "SCORE_MATCH",
    "VALIDADO",
    "VALOR_BUSCADO",
    "VALOR_ENCONTRADO",
    "FUENTE_ARCHIVO",
    "FUENTE_HOJA",
    "TIPO_BUSQUEDA",
    "OBSERVACION",
}


def get_result_sheet_names(file_path: str) -> list[str]:
    path = Path(file_path)

    if not path.exists():
        raise FileNotFoundError(f"No existe el archivo resultado: {path}")

    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        return list(workbook.sheetnames)
    finally:
        workbook.close()


def read_result_sheet_sample(file_path: str, sheet_ref: str, max_rows: int = 5, max_cols: int = 8) -> dict:
    path = Path(file_path)

    if not path.exists():
        raise FileNotFoundError(f"No existe el archivo resultado: {path}")

    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        sheet_names = workbook.sheetnames
        sheet_index, sheet_name = parse_sheet_ref(sheet_ref, sheet_names)
        ws = workbook[sheet_name]

        limit_rows = min(max_rows, ws.max_row or 0)
        limit_cols = min(max_cols, ws.max_column or 0)

        rows = []
        if limit_rows > 0 and limit_cols > 0:
            for row in ws.iter_rows(
                min_row=1,
                max_row=limit_rows,
                min_col=1,
                max_col=limit_cols,
                values_only=True,
            ):
                rows.append(["" if value is None else str(value) for value in row])

        return {
            "sheet_index": sheet_index,
            "sheet_name": sheet_name,
            "max_row": ws.max_row or 0,
            "max_column": ws.max_column or 0,
            "rows": rows,
        }
    finally:
        workbook.close()


def normalize_header_name(value) -> str:
    return str(value).strip().upper()


def read_result_headers(file_path: str, sheet_ref: str, header_row: int) -> dict:
    path = Path(file_path)

    if not path.exists():
        raise FileNotFoundError(f"No existe el archivo resultado: {path}")

    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        sheet_names = workbook.sheetnames
        sheet_index, sheet_name = parse_sheet_ref(sheet_ref, sheet_names)

        ws = workbook[sheet_name]

        if header_row < 1:
            raise ValueError("La fila header debe ser mayor o igual a 1")

        if header_row > ws.max_row:
            raise ValueError("La fila header supera el maximo de filas")

        row = next(
            ws.iter_rows(
                min_row=header_row,
                max_row=header_row,
                values_only=True,
            ),
            [],
        )

        headers = []
        for index, value in enumerate(row, start=1):
            if value is None or str(value).strip() == "":
                headers.append(f"COL_{index}")
            else:
                headers.append(str(value).strip())

        while headers and headers[-1].startswith("COL_"):
            headers.pop()

        if not headers:
            raise ValueError("No se detectaron headers en el resultado")

        missing = get_missing_required_technical_columns(headers)

        return {
            "sheet_index": sheet_index,
            "sheet_name": sheet_name,
            "header_row": header_row,
            "headers": headers,
            "first_10_headers": headers[:10],
            "first_30_headers": headers[:30],
            "max_row": ws.max_row,
            "max_column": ws.max_column,
            "missing_required_technical": missing,
        }
    finally:
        workbook.close()


def get_missing_required_technical_columns(headers: list[str]) -> list[str]:
    normalized = {normalize_header_name(h) for h in headers}
    missing = []

    for col in sorted(REQUIRED_TECHNICAL_COLUMNS):
        if col not in normalized:
            missing.append(col)

    return missing


def resolve_dataframe_column(headers: list[str], column_ref: str) -> str:
    value = str(column_ref).strip()

    if not value:
        raise ValueError("La columna no puede estar vacia")

    for header in headers:
        if value.lower() == str(header).lower():
            return header

    if value.isdigit():
        index_1 = int(value)
        if index_1 < 1 or index_1 > len(headers):
            raise ValueError(f"La columna {value} esta fuera de rango")
        return headers[index_1 - 1]

    try:
        index_1 = column_index_from_string(value.upper())
        if index_1 < 1 or index_1 > len(headers):
            raise ValueError(f"La columna {value} esta fuera de rango")
        return headers[index_1 - 1]
    except Exception:
        pass

    raise ValueError(f"No se encontro la columna: {value}")


def validate_destination_column(headers: list[str], column_ref: str) -> str:
    column_name = resolve_dataframe_column(headers, column_ref)

    if normalize_header_name(column_name) in TECHNICAL_COLUMNS:
        raise ValueError(
            f"La columna destino '{column_name}' es tecnica del sistema. "
            "Selecciona otra columna para colocar datos de la base."
        )

    return column_name


def is_validado_zero(value) -> bool:
    if value is None:
        return False

    text = str(value).strip().upper()

    if text in {"0", "0.0", "NO", "FALSE", "FALSO"}:
        return True

    return False


def build_base_index(base_records: list[dict], base_search_col: dict) -> dict:
    index = {}

    for base_record in base_records:
        raw_value = get_search_value(base_record, base_search_col)
        key = normalize_text(raw_value)

        if not key:
            continue

        if key not in index:
            index[key] = base_record

    return index


def build_base_choices(base_records: list[dict], base_search_col: dict) -> tuple[list[str], list[int]]:
    choices = []
    choice_to_record_index = []

    for index, base_record in enumerate(base_records):
        raw_value = get_search_value(base_record, base_search_col)
        normalized = normalize_text(raw_value)

        if not normalized:
            continue

        choices.append(normalized)
        choice_to_record_index.append(index)

    return choices, choice_to_record_index


def estado_to_validado(estado: str) -> int:
    if estado in {ESTADO_EXACTO, ESTADO_APROXIMADO}:
        return 1
    return 0


def update_technical_columns(
    df: pd.DataFrame,
    row_index,
    config: dict,
    estado: str,
    score,
    valor_buscado,
    valor_encontrado,
    observacion: str,
    base_record_found: bool,
):
    df.at[row_index, "ESTADO_MATCH"] = estado
    df.at[row_index, "SCORE_MATCH"] = score
    df.at[row_index, "VALIDADO"] = estado_to_validado(estado)
    df.at[row_index, "VALOR_BUSCADO"] = valor_buscado
    df.at[row_index, "VALOR_ENCONTRADO"] = valor_encontrado

    if base_record_found:
        df.at[row_index, "FUENTE_ARCHIVO"] = config["base"]["archivo"]["nombre"]
        df.at[row_index, "FUENTE_HOJA"] = config["base"]["hoja"]["nombre_detectado"]
    else:
        df.at[row_index, "FUENTE_ARCHIVO"] = ""
        df.at[row_index, "FUENTE_HOJA"] = ""

    df.at[row_index, "TIPO_BUSQUEDA"] = config["match"]["tipo_busqueda"]
    df.at[row_index, "OBSERVACION"] = observacion


def update_mapped_data_columns(
    df: pd.DataFrame,
    row_index,
    base_record: dict,
    mappings: list[dict],
):
    for mapping in mappings:
        base_alias = mapping["base_alias"]
        destination_column = mapping["destination_column_resolved"]
        df.at[row_index, destination_column] = base_record.get(base_alias)


def run_completion(config: dict) -> tuple[pd.DataFrame, dict]:
    previous = config["previous_result"]
    file_path = Path(previous["archivo"]["ruta"])
    sheet_name = previous["hoja"]["nombre_detectado"]
    header_row = int(previous["fila_header"])

    df = pd.read_excel(
        file_path,
        sheet_name=sheet_name,
        header=header_row - 1,
        dtype=object,
        engine="openpyxl",
    )

    df = df.dropna(how="all").copy()
    headers = list(df.columns)

    missing = get_missing_required_technical_columns(headers)
    if missing:
        raise ValueError(
            "El Excel no es valido para completar. Faltan columnas tecnicas: "
            + ", ".join(missing)
        )

    validado_col = resolve_dataframe_column(headers, "VALIDADO")
    search_col = resolve_dataframe_column(headers, previous["columna_busqueda"])

    mappings = []
    used_destinations = set()

    for item in config["column_mappings"]:
        destination = validate_destination_column(headers, item["destination_column_ref"])
        destination_key = normalize_header_name(destination)

        if destination_key in used_destinations:
            raise ValueError(f"Columna destino duplicada: {destination}")

        used_destinations.add(destination_key)

        mappings.append({
            "base_alias": item["base_alias"],
            "base_header_detectado": item["base_header_detectado"],
            "destination_column_ref": item["destination_column_ref"],
            "destination_column_resolved": destination,
        })

    if not mappings:
        raise ValueError("Debes configurar al menos un mapeo de columnas")

    rows_to_process = []
    for idx, value in df[validado_col].items():
        if is_validado_zero(value):
            rows_to_process.append(idx)

    base_records = read_configured_excel(config["base"], "base")

    if not base_records:
        raise ValueError("La base no tiene registros para buscar")

    base_search_col = config["base"]["columna_busqueda"]
    search_type = config["match"]["tipo_busqueda"]

    exactos = 0
    aproximados = 0
    no_encontrados = 0

    if search_type == "exacta":
        base_index = build_base_index(base_records, base_search_col)

        for idx in rows_to_process:
            valor_buscado = df.at[idx, search_col]
            key = normalize_text(valor_buscado)

            if not key:
                update_technical_columns(
                    df=df,
                    row_index=idx,
                    config=config,
                    estado=ESTADO_NO_ENCONTRADO,
                    score=0,
                    valor_buscado=valor_buscado,
                    valor_encontrado="",
                    observacion="Valor de busqueda vacio en resultado anterior",
                    base_record_found=False,
                )
                no_encontrados += 1
                continue

            base_record = base_index.get(key)

            if base_record is None:
                update_technical_columns(
                    df=df,
                    row_index=idx,
                    config=config,
                    estado=ESTADO_NO_ENCONTRADO,
                    score=0,
                    valor_buscado=valor_buscado,
                    valor_encontrado="",
                    observacion="No encontrado en nueva base",
                    base_record_found=False,
                )
                no_encontrados += 1
                continue

            valor_encontrado = get_search_value(base_record, base_search_col)

            update_mapped_data_columns(df, idx, base_record, mappings)
            update_technical_columns(
                df=df,
                row_index=idx,
                config=config,
                estado=ESTADO_EXACTO,
                score=100,
                valor_buscado=valor_buscado,
                valor_encontrado=valor_encontrado,
                observacion="Completado con coincidencia exacta",
                base_record_found=True,
            )
            exactos += 1

    elif search_type == "similaridad":
        score_exacto = int(config["match"]["score_exacto"])
        score_aproximado = int(config["match"]["score_aproximado"])

        choices, choice_to_record_index = build_base_choices(base_records, base_search_col)

        if not choices:
            raise ValueError("La base no tiene valores validos para busqueda por similaridad")

        for idx in rows_to_process:
            valor_buscado = df.at[idx, search_col]
            query = normalize_text(valor_buscado)

            if not query:
                update_technical_columns(
                    df=df,
                    row_index=idx,
                    config=config,
                    estado=ESTADO_NO_ENCONTRADO,
                    score=0,
                    valor_buscado=valor_buscado,
                    valor_encontrado="",
                    observacion="Valor de busqueda vacio en resultado anterior",
                    base_record_found=False,
                )
                no_encontrados += 1
                continue

            match = find_best_similarity(
                query,
                choices,
            )

            if match is None:
                update_technical_columns(
                    df=df,
                    row_index=idx,
                    config=config,
                    estado=ESTADO_NO_ENCONTRADO,
                    score=0,
                    valor_buscado=valor_buscado,
                    valor_encontrado="",
                    observacion="No encontrado en nueva base",
                    base_record_found=False,
                )
                no_encontrados += 1
                continue

            matched_text, score, choice_index = match
            score = round(float(score), 2)

            base_record_index = choice_to_record_index[choice_index]
            base_record = base_records[base_record_index]
            valor_encontrado = get_search_value(base_record, base_search_col)

            if score >= score_exacto:
                estado = ESTADO_EXACTO
                observacion = "Completado por similaridad en rango EXACTO"
                exactos += 1
            elif score >= score_aproximado:
                estado = ESTADO_APROXIMADO
                observacion = "Completado por similaridad en rango APROXIMADO"
                aproximados += 1
            else:
                estado = ESTADO_NO_ENCONTRADO
                base_record = None
                valor_encontrado = ""
                observacion = "Sigue sin encontrarse en nueva base"
                no_encontrados += 1

            if base_record is not None:
                update_mapped_data_columns(df, idx, base_record, mappings)

            update_technical_columns(
                df=df,
                row_index=idx,
                config=config,
                estado=estado,
                score=score,
                valor_buscado=valor_buscado,
                valor_encontrado=valor_encontrado,
                observacion=observacion,
                base_record_found=base_record is not None,
            )

    else:
        raise ValueError(f"Tipo de busqueda no soportado: {search_type}")

    summary = {
        "total_filas_resultado": len(df),
        "filas_reprocesadas": len(rows_to_process),
        "base_rows": len(base_records),
        "exactos_nuevos": exactos,
        "aproximados_nuevos": aproximados,
        "no_encontrados_finales": no_encontrados,
    }

    return df, summary


def unique_sheet_name(workbook, wanted_name: str) -> str:
    base = sanitize_sheet_name(wanted_name)
    if base not in workbook.sheetnames:
        return base

    counter = 2
    while True:
        suffix = f"_{counter}"
        max_len = 31 - len(suffix)
        candidate = f"{base[:max_len]}{suffix}"

        if candidate not in workbook.sheetnames:
            return candidate

        counter += 1


def write_dataframe_to_openpyxl_sheet(workbook, sheet_name: str, df: pd.DataFrame):
    ws = workbook.create_sheet(sheet_name)

    header_fill = PatternFill("solid", fgColor="5E6F32")
    header_font = Font(color="FFFFFF", bold=True)

    for col_idx, col_name in enumerate(df.columns, start=1):
        cell = ws.cell(row=1, column=col_idx)
        cell.value = col_name
        cell.fill = header_fill
        cell.font = header_font

    for row_idx, row_values in enumerate(df.itertuples(index=False, name=None), start=2):
        for col_idx, value in enumerate(row_values, start=1):
            ws.cell(row=row_idx, column=col_idx).value = value

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions

    for column_cells in ws.columns:
        max_length = 12
        column_letter = column_cells[0].column_letter

        for cell in column_cells[:100]:
            value = cell.value
            if value is not None:
                max_length = max(max_length, len(str(value)) + 2)

        ws.column_dimensions[column_letter].width = min(max_length, 45)


def export_completion_result(df: pd.DataFrame, config: dict) -> Path:
    previous = config["previous_result"]
    original_path = Path(previous["archivo"]["ruta"])
    output_mode = config["completion_output"]["modo"]
    wanted_sheet = config["completion_output"]["nombre_hoja"]

    if output_mode == "excel_nuevo":
        output_dir = original_path.parent
        timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        output_path = output_dir / f"{original_path.stem}_completado_{timestamp}.xlsx"
        sheet_name = sanitize_sheet_name(wanted_sheet)

        with pd.ExcelWriter(output_path, engine="xlsxwriter") as writer:
            df.to_excel(writer, index=False, sheet_name=sheet_name)
            apply_xlsxwriter_format(writer, df, sheet_name)

        return output_path

    if output_mode == "hoja_nueva_mismo_excel":
        workbook = load_workbook(original_path)
        sheet_name = unique_sheet_name(workbook, wanted_sheet)
        write_dataframe_to_openpyxl_sheet(workbook, sheet_name, df)
        workbook.save(original_path)
        workbook.close()
        return original_path

    raise ValueError(f"Modo de salida de completado no soportado: {output_mode}")
