from pathlib import Path
from datetime import datetime
import re
from copy import deepcopy
from uuid import uuid4

import pandas as pd
from openpyxl.styles import PatternFill, Font
from openpyxl.utils.cell import column_index_from_string
from src.core.excel_safety import (
    captured_sources, attach_capture, atomic_write, atomic_update, ExcelHeadersChangedError,
)

from src.core.search_executor import normalize_text, sanitize_sheet_name, apply_xlsxwriter_format


DRIVER_PORCENTAJE = "porcentaje"
DRIVER_CANTIDAD = "cantidad"

COL_FILA_ORIGEN = "FILA_ORIGEN"
COL_DNI = "DNI"
COL_CECO = "CECO"
COL_CECO2 = "CECO2"
COL_PORCENTAJE = "PORCENTAJE"
COL_PRECIO = "PRECIO"
COL_PRECIO_DISTRIBUIDO = "PRECIO_DISTRIBUIDO"
COL_OBSERVACION = "OBSERVACION"
COL_GRUPO_FILTRO = "GRUPO_FILTRO"
COL_VALOR_FILTRO = "VALOR_FILTRO"
COL_CANTIDAD_EQUIVALENTE = "CANTIDAD_EQUIVALENTE"
COL_PORCENTAJE_DISTRIBUCION = "PORCENTAJE_DISTRIBUCION"

COL_VALIDADO = "VALIDADO"


def resolve_dataframe_column(headers: list[str], column_ref: str) -> str:
    value = str(column_ref).strip()

    if not value:
        raise ValueError("La columna no puede estar vacia")

    for header in headers:
        if value.lower() == str(header).lower():
            return header

    if value.isdigit():
        index_1 = int(value)
        if index_1 < 1 or index_1 > len(headers):
            raise ValueError(f"La columna {value} esta fuera de rango")
        return headers[index_1 - 1]

    try:
        index_1 = column_index_from_string(value.upper())
        if index_1 < 1 or index_1 > len(headers):
            raise ValueError(f"La columna {value} esta fuera de rango")
        return headers[index_1 - 1]
    except Exception:
        pass

    raise ValueError(f"No se encontro la columna: {value}")


def read_excel_with_origin(file_path: str, sheet_name: str, header_row: int) -> pd.DataFrame:
    path = Path(file_path)

    if not path.exists():
        raise FileNotFoundError(f"No existe el archivo: {path}")

    if header_row < 1:
        raise ValueError("La fila header debe ser mayor o igual a 1")

    df = pd.read_excel(
        path,
        sheet_name=sheet_name,
        header=header_row - 1,
        dtype=object,
        engine="openpyxl",
    )

    df[COL_FILA_ORIGEN] = [header_row + 1 + i for i in range(len(df))]
    df = df.dropna(how="all").copy()

    return df


def clean_cell_text(value) -> str:
    if value is None:
        return ""

    try:
        if pd.isna(value):
            return ""
    except Exception:
        pass

    text = str(value).strip()

    if text.endswith(".0"):
        possible_number = text[:-2]
        if possible_number.isdigit():
            text = possible_number

    return text.strip()


def normalize_dni(value) -> str:
    text = clean_cell_text(value)

    if not text:
        return ""

    # Caso comun Excel: 9344630.0 -> 9344630
    if text.endswith(".0"):
        possible_number = text[:-2]
        if possible_number.isdigit():
            text = possible_number

    # Para DNI/RUC/documentos numericos: quitar espacios, comas, puntos, guiones, etc.
    digits = re.sub(r"\D+", "", text)

    if digits:
        # Evita que 009344630 y 9344630 no coincidan.
        # Si todo queda en ceros, conserva un solo 0.
        digits = digits.lstrip("0") or "0"
        return digits

    # Fallback por si algun documento tiene letras.
    return normalize_text(text)


def normalize_ceco(value) -> str:
    text = clean_cell_text(value)
    text = text.upper()
    text = re.sub(r"\s+", "", text)
    return text


def is_validado_one(value) -> bool:
    text = clean_cell_text(value).upper()
    return text in {"1", "1.0", "SI", "TRUE", "VERDADERO"}


def is_empty_value(value) -> bool:
    return clean_cell_text(value) == ""


def parse_number(value):
    if value is None:
        return None

    try:
        if pd.isna(value):
            return None
    except Exception:
        pass

    if isinstance(value, (int, float)):
        return float(value)

    text = str(value).strip()

    if not text:
        return None

    text = text.replace("S/", "")
    text = text.replace("$", "")
    text = text.replace("USD", "")
    text = text.replace("PEN", "")
    text = text.strip()

    if "," in text and "." in text:
        text = text.replace(",", "")
    elif "," in text and "." not in text:
        text = text.replace(",", ".")

    try:
        return float(text)
    except Exception:
        return None


def parse_percentage(value):
    number = parse_number(value)

    if number is None:
        return None

    if number > 1:
        number = number / 100.0

    return number


def apply_special_ceco_rules(ceco_value, rules: list[dict]) -> tuple[bool, str, str]:
    ceco = normalize_ceco(ceco_value)

    if not ceco:
        return False, "", ""

    for rule in rules:
        prefijo = normalize_ceco(rule.get("prefijo", ""))
        reemplazo = normalize_ceco(rule.get("ceco2", ""))

        if not prefijo or not reemplazo:
            continue

        if ceco.startswith(prefijo):
            obs = f"CECO especial por prefijo {prefijo}"
            return True, reemplazo, obs

    return False, ceco, ""


def build_base_index_by_dni(
    base_df: pd.DataFrame,
    base_dni_col: str,
    base_ceco_col: str,
    base_percentage_col: str,
) -> dict:
    index = {}

    for _, row in base_df.iterrows():
        dni_key = normalize_dni(row.get(base_dni_col))

        if not dni_key:
            continue

        ceco = row.get(base_ceco_col)
        porcentaje = row.get(base_percentage_col)

        item = {
            "dni": clean_cell_text(row.get(base_dni_col)),
            "ceco": clean_cell_text(ceco),
            "porcentaje_raw": porcentaje,
            "porcentaje": parse_percentage(porcentaje),
        }

        index.setdefault(dni_key, []).append(item)

    return index


def normalize_filter_value(value) -> str:
    return normalize_text(clean_cell_text(value))


def prepare_filter_groups(filter_cfg: dict) -> list[dict]:
    groups = []

    if not filter_cfg or not filter_cfg.get("enabled"):
        return groups

    for group in filter_cfg.get("groups", []):
        name = clean_cell_text(group.get("name"))

        raw_values = group.get("values", [])
        values = []
        normalized_values = set()

        for value in raw_values:
            clean_value = clean_cell_text(value)

            if not clean_value:
                continue

            values.append(clean_value)
            normalized_values.add(normalize_filter_value(clean_value))

        if not name or not values:
            continue

        groups.append({
            "name": name,
            "values": values,
            "normalized_values": normalized_values,
        })

    return groups


def get_filter_matches(row, filter_col: str | None, filter_groups: list[dict]) -> list[tuple[str | None, str | None]]:
    if not filter_col:
        return [(None, None)]

    raw_value = clean_cell_text(row.get(filter_col))
    normalized_value = normalize_filter_value(raw_value)

    matches = []

    for group in filter_groups:
        if normalized_value in group["normalized_values"]:
            matches.append((group["name"], raw_value))

    return matches




PROTECTED_DRIVER_OUTPUT_COLUMNS = {
    COL_GRUPO_FILTRO,
    COL_VALOR_FILTRO,
    COL_FILA_ORIGEN,
    COL_DNI,
    COL_CECO,
    COL_CECO2,
    COL_PORCENTAJE,
    COL_PRECIO,
    COL_PRECIO_DISTRIBUIDO,
    COL_OBSERVACION,
    COL_VALIDADO,
}


def make_unique_output_name(name: str, existing: set[str]) -> str:
    base = clean_cell_text(name) or "EXTRA"

    if base in PROTECTED_DRIVER_OUTPUT_COLUMNS:
        base = f"ORIGEN_{base}"

    candidate = base
    counter = 2

    while candidate in existing:
        candidate = f"{base}_{counter}"
        counter += 1

    existing.add(candidate)
    return candidate


def resolve_extra_result_columns(
    result_headers: list[str],
    result_cfg: dict,
    used_result_cols: list[str | None],
) -> list[dict]:
    extra_cols_cfg = result_cfg.get("extra_cols", []) or []
    used_headers = {clean_cell_text(value) for value in used_result_cols if clean_cell_text(value)}
    existing_output_names = set(PROTECTED_DRIVER_OUTPUT_COLUMNS)
    resolved = []

    for item in extra_cols_cfg:
        raw_header = item.get("header") if isinstance(item, dict) else item
        raw_alias = item.get("alias") if isinstance(item, dict) else raw_header

        if not clean_cell_text(raw_header):
            continue

        header = resolve_dataframe_column(result_headers, raw_header)

        if clean_cell_text(header) in used_headers:
            continue

        output_name = make_unique_output_name(raw_alias or header, existing_output_names)

        resolved.append({
            "header": header,
            "output_name": output_name,
        })

    return resolved


def build_extra_values_from_row(row, extra_result_cols: list[dict]) -> dict:
    values = {}

    for item in extra_result_cols:
        output_name = item["output_name"]
        header = item["header"]
        values[output_name] = row.get(header)

    return values

def build_driver_row(
    fila_origen,
    dni,
    ceco,
    ceco2,
    porcentaje,
    observacion,
    driver_type,
    precio=None,
    grupo_filtro=None,
    valor_filtro=None,
    extra_values=None,
):
    row = {}

    if grupo_filtro is not None:
        row[COL_GRUPO_FILTRO] = clean_cell_text(grupo_filtro)
        row[COL_VALOR_FILTRO] = clean_cell_text(valor_filtro)

    if extra_values:
        for key, value in extra_values.items():
            key = clean_cell_text(key)
            if key and key not in row:
                row[key] = value

    row.update({
        COL_FILA_ORIGEN: fila_origen,
        COL_DNI: clean_cell_text(dni),
        COL_CECO: clean_cell_text(ceco),
        COL_CECO2: clean_cell_text(ceco2),
        COL_PORCENTAJE: porcentaje,
    })

    if driver_type == DRIVER_CANTIDAD:
        precio_number = parse_number(precio)

        row[COL_PRECIO] = precio_number

        if precio_number is None or porcentaje is None:
            row[COL_PRECIO_DISTRIBUIDO] = None
        else:
            row[COL_PRECIO_DISTRIBUIDO] = precio_number * porcentaje

    row[COL_OBSERVACION] = observacion

    return row

def build_observed_row(fila_origen, dni, ceco, motivo, grupo_filtro=None, valor_filtro=None):
    row = {}

    if grupo_filtro is not None:
        row[COL_GRUPO_FILTRO] = clean_cell_text(grupo_filtro)
        row[COL_VALOR_FILTRO] = clean_cell_text(valor_filtro)

    row.update({
        COL_FILA_ORIGEN: fila_origen,
        COL_DNI: clean_cell_text(dni),
        COL_CECO: clean_cell_text(ceco),
        "MOTIVO": motivo,
    })

    return row


def validate_driver_config(config: dict):
    driver_type = config.get("driver_type")

    if driver_type not in {DRIVER_PORCENTAJE, DRIVER_CANTIDAD}:
        raise ValueError("driver_type debe ser 'porcentaje' o 'cantidad'")

    result = config.get("resultado", {})
    base = config.get("base", {})
    filter_cfg = config.get("filter", {}) or {}

    required_result = ["file_path", "sheet_name", "header_row", "dni_col", "ceco_col"]
    required_base = ["file_path", "sheet_name", "header_row", "dni_col", "ceco_col", "percentage_col"]

    for key in required_result:
        if key not in result or str(result.get(key)).strip() == "":
            raise ValueError(f"Falta resultado.{key}")

    if driver_type == DRIVER_CANTIDAD:
        if "price_col" not in result or str(result.get("price_col")).strip() == "":
            raise ValueError("Para driver por cantidad falta resultado.price_col")

    for key in required_base:
        if key not in base or str(base.get(key)).strip() == "":
            raise ValueError(f"Falta base.{key}")

    if filter_cfg.get("enabled"):
        if "filter_col" not in result or str(result.get("filter_col")).strip() == "":
            raise ValueError("Para usar filtro falta resultado.filter_col")

        groups = filter_cfg.get("groups", [])

        if not groups:
            raise ValueError("Para usar filtro debes configurar al menos un grupo")

        for index, group in enumerate(groups, start=1):
            name = clean_cell_text(group.get("name"))
            values = [clean_cell_text(value) for value in group.get("values", []) if clean_cell_text(value)]

            if not name:
                raise ValueError(f"El grupo de filtro #{index} no tiene nombre")

            if not values:
                raise ValueError(f"El grupo de filtro '{name}' no tiene valores seleccionados")


def run_driver(config: dict, progress=None) -> tuple[pd.DataFrame, pd.DataFrame, dict]:
    validate_driver_config(config)
    if progress:
        progress("Capturando resultado y maestro...")
    with captured_sources({role: config[role]["file_path"] for role in ("resultado", "base")}) as captures:
        local = deepcopy(config)
        if progress:
            progress("Validando encabezados...")
        for role in captures:
            local[role]["file_path"] = captures[role].snapshot
            expected_headers = local[role].get("expected_headers")
            if expected_headers is not None:
                from src.core.excel_inspector import read_headers
                actual = read_headers(local[role]["file_path"], local[role]["sheet_name"], int(local[role]["header_row"]))["headers"]
                if actual != expected_headers:
                    raise ExcelHeadersChangedError("Los encabezados cambiaron. Vuelva a leerlos antes de generar el driver.")
        if progress:
            progress("Generando drivers...")
        df, observed, summary = _run_driver(local)
        attach_capture(df, captures, "resultado")
        return df, observed, summary


def _run_driver(config: dict) -> tuple[pd.DataFrame, pd.DataFrame, dict]:
    validate_driver_config(config)

    driver_type = config["driver_type"]
    result_cfg = config["resultado"]
    base_cfg = config["base"]
    special_rules = config.get("special_ceco_rules", [])
    filter_cfg = config.get("filter", {}) or {}
    filter_enabled = bool(filter_cfg.get("enabled"))
    filter_groups = prepare_filter_groups(filter_cfg)

    result_df = read_excel_with_origin(
        file_path=result_cfg["file_path"],
        sheet_name=result_cfg["sheet_name"],
        header_row=int(result_cfg["header_row"]),
    )

    base_df = read_excel_with_origin(
        file_path=base_cfg["file_path"],
        sheet_name=base_cfg["sheet_name"],
        header_row=int(base_cfg["header_row"]),
    )

    result_headers = list(result_df.columns)
    base_headers = list(base_df.columns)

    result_dni_col = resolve_dataframe_column(result_headers, result_cfg["dni_col"])
    result_ceco_col = resolve_dataframe_column(result_headers, result_cfg["ceco_col"])

    result_price_col = None
    if driver_type == DRIVER_CANTIDAD:
        result_price_col = resolve_dataframe_column(result_headers, result_cfg["price_col"])

    result_filter_col = None
    if filter_enabled:
        result_filter_col = resolve_dataframe_column(result_headers, result_cfg["filter_col"])

    extra_result_cols = resolve_extra_result_columns(
        result_headers=result_headers,
        result_cfg=result_cfg,
        used_result_cols=[result_dni_col, result_ceco_col, result_price_col, result_filter_col, COL_VALIDADO],
    )

    validado_col = None
    if COL_VALIDADO in result_headers:
        validado_col = COL_VALIDADO

    base_dni_col = resolve_dataframe_column(base_headers, base_cfg["dni_col"])
    base_ceco_col = resolve_dataframe_column(base_headers, base_cfg["ceco_col"])
    base_percentage_col = resolve_dataframe_column(base_headers, base_cfg["percentage_col"])

    base_index = build_base_index_by_dni(
        base_df=base_df,
        base_dni_col=base_dni_col,
        base_ceco_col=base_ceco_col,
        base_percentage_col=base_percentage_col,
    )

    driver_rows = []
    observed_rows = []

    skipped_not_validated = 0
    special_count = 0
    empty_dni_count = 0
    not_found_count = 0
    expanded_count = 0
    invalid_percentage_count = 0
    invalid_price_count = 0
    excluded_by_filter_count = 0

    for _, row in result_df.iterrows():
        filter_matches = get_filter_matches(row, result_filter_col, filter_groups)

        if filter_enabled and not filter_matches:
            excluded_by_filter_count += 1
            continue

        extra_values = build_extra_values_from_row(row, extra_result_cols)

        for grupo_filtro, valor_filtro in filter_matches:
            fila_origen = row.get(COL_FILA_ORIGEN)
            dni = row.get(result_dni_col)
            ceco_original = row.get(result_ceco_col)

            if validado_col is not None and not is_validado_one(row.get(validado_col)):
                skipped_not_validated += 1
                observed_rows.append(
                    build_observed_row(
                        fila_origen=fila_origen,
                        dni=dni,
                        ceco=ceco_original,
                        motivo="Fila no procesada porque VALIDADO no es 1",
                        grupo_filtro=grupo_filtro,
                        valor_filtro=valor_filtro,
                    )
                )
                continue

            precio = None
            if driver_type == DRIVER_CANTIDAD:
                precio = row.get(result_price_col)
                if parse_number(precio) is None:
                    invalid_price_count += 1
                    observed_rows.append(
                        build_observed_row(
                            fila_origen=fila_origen,
                            dni=dni,
                            ceco=ceco_original,
                            motivo="Precio vacio o no numerico",
                            grupo_filtro=grupo_filtro,
                            valor_filtro=valor_filtro,
                        )
                    )

            is_special, ceco2_special, special_obs = apply_special_ceco_rules(
                ceco_original,
                special_rules,
            )

            if is_special:
                special_count += 1
                driver_rows.append(
                    build_driver_row(
                        fila_origen=fila_origen,
                        dni=dni,
                        ceco=ceco_original,
                        ceco2=ceco2_special,
                        porcentaje=1,
                        observacion=special_obs,
                        driver_type=driver_type,
                        precio=precio,
                        grupo_filtro=grupo_filtro,
                        valor_filtro=valor_filtro,
                        extra_values=extra_values,
                    )
                )
                continue

            dni_key = normalize_dni(dni)

            if not dni_key:
                empty_dni_count += 1
                driver_rows.append(
                    build_driver_row(
                        fila_origen=fila_origen,
                        dni=dni,
                        ceco=ceco_original,
                        ceco2=ceco_original,
                        porcentaje=1,
                        observacion="DNI vacio, se usa CECO original con porcentaje 1",
                        driver_type=driver_type,
                        precio=precio,
                        grupo_filtro=grupo_filtro,
                        valor_filtro=valor_filtro,
                        extra_values=extra_values,
                    )
                )
                continue

            base_matches = base_index.get(dni_key, [])

            if not base_matches:
                not_found_count += 1
                driver_rows.append(
                    build_driver_row(
                        fila_origen=fila_origen,
                        dni=dni,
                        ceco=ceco_original,
                        ceco2=ceco_original,
                        porcentaje=1,
                        observacion="DNI no encontrado en base, se usa CECO original con porcentaje 1",
                        driver_type=driver_type,
                        precio=precio,
                        grupo_filtro=grupo_filtro,
                        valor_filtro=valor_filtro,
                        extra_values=extra_values,
                    )
                )
                continue

            for match in base_matches:
                porcentaje = match["porcentaje"]

                if porcentaje is None:
                    invalid_percentage_count += 1
                    porcentaje = 1
                    obs = "DNI encontrado, porcentaje vacio o no numerico; se usa porcentaje 1"
                else:
                    obs = "DNI encontrado en base, se expande CECO y porcentaje"

                driver_rows.append(
                    build_driver_row(
                        fila_origen=fila_origen,
                        dni=dni,
                        ceco=ceco_original,
                        ceco2=match["ceco"],
                        porcentaje=porcentaje,
                        observacion=obs,
                        driver_type=driver_type,
                        precio=precio,
                        grupo_filtro=grupo_filtro,
                        valor_filtro=valor_filtro,
                        extra_values=extra_values,
                    )
                )
                expanded_count += 1

    df_driver = pd.DataFrame(driver_rows)
    df_observados = pd.DataFrame(observed_rows)

    if df_driver.empty:
        driver_columns = []

        if filter_enabled:
            driver_columns.extend([COL_GRUPO_FILTRO, COL_VALOR_FILTRO])

        driver_columns.extend([item["output_name"] for item in extra_result_cols])

        driver_columns.extend([
            COL_FILA_ORIGEN,
            COL_DNI,
            COL_CECO,
            COL_CECO2,
            COL_PORCENTAJE,
        ])

        if driver_type == DRIVER_CANTIDAD:
            driver_columns.extend([COL_PRECIO, COL_PRECIO_DISTRIBUIDO])

        driver_columns.append(COL_OBSERVACION)
        df_driver = pd.DataFrame(columns=driver_columns)

    if df_observados.empty:
        observed_columns = []

        if filter_enabled:
            observed_columns.extend([COL_GRUPO_FILTRO, COL_VALOR_FILTRO])

        observed_columns.extend([COL_FILA_ORIGEN, COL_DNI, COL_CECO, "MOTIVO"])
        df_observados = pd.DataFrame(columns=observed_columns)

    summary = {
        "tipo_driver": driver_type,
        "filtro_usado": "SI" if filter_enabled else "NO",
        "grupos_filtro": len(filter_groups) if filter_enabled else 0,
        "filas_excluidas_por_filtro": excluded_by_filter_count,
        "columnas_extra_copiadas": len(extra_result_cols),
        "filas_resultado": len(result_df),
        "filas_base": len(base_df),
        "filas_driver": len(df_driver),
        "filas_observadas": len(df_observados),
        "filas_no_validadas": skipped_not_validated,
        "cecos_especiales": special_count,
        "dni_vacios": empty_dni_count,
        "dni_no_encontrados": not_found_count,
        "filas_expandidas_por_dni": expanded_count,
        "porcentajes_invalidos": invalid_percentage_count,
        "precios_invalidos": invalid_price_count,
    }

    if extra_result_cols:
        summary["detalle_columnas_extra"] = " | ".join(item["output_name"] for item in extra_result_cols)

    if filter_enabled:
        summary["detalle_grupos_filtro"] = " | ".join(
            f"{group['name']}: {', '.join(group['values'])}"
            for group in filter_groups
        )

    return df_driver, df_observados, summary

def build_driver_debug_report(
    config: dict,
    output_dir: str | Path | None = None,
    max_rows: int = 500,
) -> Path:
    validate_driver_config(config)

    driver_type = config["driver_type"]
    result_cfg = config["resultado"]
    base_cfg = config["base"]

    result_df = read_excel_with_origin(
        file_path=result_cfg["file_path"],
        sheet_name=result_cfg["sheet_name"],
        header_row=int(result_cfg["header_row"]),
    )

    base_df = read_excel_with_origin(
        file_path=base_cfg["file_path"],
        sheet_name=base_cfg["sheet_name"],
        header_row=int(base_cfg["header_row"]),
    )

    result_headers = list(result_df.columns)
    base_headers = list(base_df.columns)

    result_dni_col = resolve_dataframe_column(result_headers, result_cfg["dni_col"])
    result_ceco_col = resolve_dataframe_column(result_headers, result_cfg["ceco_col"])

    result_price_col = None
    if driver_type == DRIVER_CANTIDAD:
        result_price_col = resolve_dataframe_column(result_headers, result_cfg["price_col"])

    validado_col = COL_VALIDADO if COL_VALIDADO in result_headers else None

    base_dni_col = resolve_dataframe_column(base_headers, base_cfg["dni_col"])
    base_ceco_col = resolve_dataframe_column(base_headers, base_cfg["ceco_col"])
    base_percentage_col = resolve_dataframe_column(base_headers, base_cfg["percentage_col"])

    base_index = build_base_index_by_dni(
        base_df=base_df,
        base_dni_col=base_dni_col,
        base_ceco_col=base_ceco_col,
        base_percentage_col=base_percentage_col,
    )

    config_rows = [
        {"CAMPO": "driver_type", "VALOR": driver_type},
        {"CAMPO": "resultado_file", "VALOR": result_cfg["file_path"]},
        {"CAMPO": "resultado_sheet", "VALOR": result_cfg["sheet_name"]},
        {"CAMPO": "resultado_header_row", "VALOR": result_cfg["header_row"]},
        {"CAMPO": "resultado_dni_col_resuelta", "VALOR": result_dni_col},
        {"CAMPO": "resultado_ceco_col_resuelta", "VALOR": result_ceco_col},
        {"CAMPO": "resultado_price_col_resuelta", "VALOR": result_price_col or ""},
        {"CAMPO": "resultado_validado_col", "VALOR": validado_col or "NO_EXISTE"},
        {"CAMPO": "resultado_filas_leidas", "VALOR": len(result_df)},
        {"CAMPO": "base_file", "VALOR": base_cfg["file_path"]},
        {"CAMPO": "base_sheet", "VALOR": base_cfg["sheet_name"]},
        {"CAMPO": "base_header_row", "VALOR": base_cfg["header_row"]},
        {"CAMPO": "base_dni_col_resuelta", "VALOR": base_dni_col},
        {"CAMPO": "base_ceco_col_resuelta", "VALOR": base_ceco_col},
        {"CAMPO": "base_percentage_col_resuelta", "VALOR": base_percentage_col},
        {"CAMPO": "base_filas_leidas", "VALOR": len(base_df)},
        {"CAMPO": "base_dnis_unicos_normalizados", "VALOR": len(base_index)},
    ]

    df_config = pd.DataFrame(config_rows)

    base_sample_rows = []
    for _, row in base_df.head(max_rows).iterrows():
        dni_raw = row.get(base_dni_col)
        dni_norm = normalize_dni(dni_raw)

        base_sample_rows.append({
            "FILA_ORIGEN_BASE": row.get(COL_FILA_ORIGEN),
            "DNI_RAW_BASE": clean_cell_text(dni_raw),
            "DNI_NORMALIZADO_BASE": dni_norm,
            "CECO_BASE": clean_cell_text(row.get(base_ceco_col)),
            "PORCENTAJE_RAW_BASE": clean_cell_text(row.get(base_percentage_col)),
            "PORCENTAJE_PARSEADO_BASE": parse_percentage(row.get(base_percentage_col)),
        })

    df_base_sample = pd.DataFrame(base_sample_rows)

    result_check_rows = []
    for _, row in result_df.head(max_rows).iterrows():
        dni_raw = row.get(result_dni_col)
        dni_norm = normalize_dni(dni_raw)
        matches = base_index.get(dni_norm, [])

        item = {
            "FILA_ORIGEN_RESULTADO": row.get(COL_FILA_ORIGEN),
            "VALIDADO": clean_cell_text(row.get(validado_col)) if validado_col else "NO_EXISTE",
            "DNI_RAW_RESULTADO": clean_cell_text(dni_raw),
            "DNI_NORMALIZADO_RESULTADO": dni_norm,
            "CECO_RESULTADO": clean_cell_text(row.get(result_ceco_col)),
            "EXISTE_EN_BASE": "SI" if matches else "NO",
            "CANTIDAD_MATCHES_BASE": len(matches),
            "CECOS_ENCONTRADOS_BASE": " | ".join(clean_cell_text(m.get("ceco")) for m in matches[:20]),
            "PORCENTAJES_ENCONTRADOS_BASE": " | ".join(clean_cell_text(m.get("porcentaje_raw")) for m in matches[:20]),
        }

        if driver_type == DRIVER_CANTIDAD:
            item["PRECIO_RESULTADO"] = clean_cell_text(row.get(result_price_col))
            item["PRECIO_PARSEADO"] = parse_number(row.get(result_price_col))

        result_check_rows.append(item)

    df_result_check = pd.DataFrame(result_check_rows)

    base_key_rows = []
    for key in list(base_index.keys())[:max_rows]:
        matches = base_index[key]
        base_key_rows.append({
            "DNI_NORMALIZADO_BASE": key,
            "CANTIDAD_FILAS": len(matches),
            "CECOS": " | ".join(clean_cell_text(m.get("ceco")) for m in matches[:20]),
            "PORCENTAJES": " | ".join(clean_cell_text(m.get("porcentaje_raw")) for m in matches[:20]),
        })

    df_base_keys = pd.DataFrame(base_key_rows)

    if output_dir is None:
        output_dir = Path("outputs") / "drivers" / "_debug_driver"
    else:
        output_dir = Path(output_dir)

    output_dir.mkdir(parents=True, exist_ok=True)

    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    output_path = output_dir / f"debug_driver_{timestamp}_{uuid4().hex[:8]}.xlsx"

    def serialize(temporary):
        with pd.ExcelWriter(temporary, engine="xlsxwriter") as writer:
            write_sheets(writer)

    def write_sheets(writer):
        sheets = [
            ("DEBUG_CONFIG", df_config),
            ("BASE_DNI_SAMPLE", df_base_sample),
            ("RESULT_DNI_CHECK", df_result_check),
            ("BASE_KEYS_SAMPLE", df_base_keys),
        ]

        for sheet_name, df in sheets:
            clean_sheet = sanitize_sheet_name(sheet_name)
            df.to_excel(writer, index=False, sheet_name=clean_sheet)
            apply_xlsxwriter_format(writer, df, clean_sheet)

    return atomic_write(output_path, serialize)




def build_ceco_summary(df_driver: pd.DataFrame, driver_type: str) -> pd.DataFrame:
    if df_driver is None or df_driver.empty:
        return pd.DataFrame(columns=[COL_CECO2, "CANTIDAD_FILAS", "IMPORTE_TOTAL"])

    if COL_CECO2 not in df_driver.columns:
        return pd.DataFrame(columns=[COL_CECO2, "CANTIDAD_FILAS", "IMPORTE_TOTAL"])

    work = df_driver.copy()
    work[COL_CECO2] = work[COL_CECO2].apply(clean_cell_text)

    group_cols = []
    has_filter = COL_GRUPO_FILTRO in work.columns

    if has_filter:
        work[COL_GRUPO_FILTRO] = work[COL_GRUPO_FILTRO].apply(clean_cell_text)
        group_cols.append(COL_GRUPO_FILTRO)

    group_cols.append(COL_CECO2)

    if driver_type == DRIVER_CANTIDAD and COL_PRECIO_DISTRIBUIDO in work.columns:
        work["_IMPORTE_TOTAL"] = pd.to_numeric(
            work[COL_PRECIO_DISTRIBUIDO],
            errors="coerce",
        ).fillna(0)

        summary_df = (
            work.groupby(group_cols, dropna=False)
            .agg(
                CANTIDAD_FILAS=(COL_CECO2, "size"),
                IMPORTE_TOTAL=("_IMPORTE_TOTAL", "sum"),
            )
            .reset_index()
        )

        summary_df = summary_df.sort_values(group_cols).reset_index(drop=True)

        total_data = {
            COL_CECO2: "TOTAL_GENERAL",
            "CANTIDAD_FILAS": int(summary_df["CANTIDAD_FILAS"].sum()),
            "IMPORTE_TOTAL": float(summary_df["IMPORTE_TOTAL"].sum()),
        }

        if COL_GRUPO_FILTRO in summary_df.columns:
            total_data[COL_GRUPO_FILTRO] = "TOTAL_GENERAL"

        total_row = pd.DataFrame([total_data])
        summary_df = pd.concat([summary_df, total_row], ignore_index=True)
        return summary_df

    # Driver por porcentaje:
    # - PORCENTAJE es la proporcion original de cada fila/persona.
    # - Al agrupar por CECO, la suma ya no debe mostrarse como 3300%.
    # - Esa suma representa cantidad equivalente: 1 + 1 + 0.35 + 0.65, etc.
    # - PORCENTAJE_DISTRIBUCION es cantidad_equivalente / total del grupo.
    work["_CANTIDAD_EQUIVALENTE"] = pd.to_numeric(
        work[COL_PORCENTAJE],
        errors="coerce",
    ).fillna(0)

    summary_df = (
        work.groupby(group_cols, dropna=False)
        .agg(
            CANTIDAD_FILAS=(COL_CECO2, "size"),
            **{COL_CANTIDAD_EQUIVALENTE: ("_CANTIDAD_EQUIVALENTE", "sum")},
        )
        .reset_index()
    )

    if has_filter:
        totals = summary_df.groupby(COL_GRUPO_FILTRO)[COL_CANTIDAD_EQUIVALENTE].transform("sum")
    else:
        total_value = float(summary_df[COL_CANTIDAD_EQUIVALENTE].sum())
        totals = pd.Series([total_value] * len(summary_df), index=summary_df.index)

    summary_df[COL_PORCENTAJE_DISTRIBUCION] = [
        (float(value) / float(total)) if float(total or 0) != 0 else 0
        for value, total in zip(summary_df[COL_CANTIDAD_EQUIVALENTE], totals)
    ]

    summary_df = summary_df.sort_values(group_cols).reset_index(drop=True)

    total_rows = []

    if has_filter:
        for group_name, group_df in summary_df.groupby(COL_GRUPO_FILTRO, dropna=False):
            total_rows.append({
                COL_GRUPO_FILTRO: group_name,
                COL_CECO2: "TOTAL_GRUPO",
                "CANTIDAD_FILAS": int(group_df["CANTIDAD_FILAS"].sum()),
                COL_CANTIDAD_EQUIVALENTE: float(group_df[COL_CANTIDAD_EQUIVALENTE].sum()),
                COL_PORCENTAJE_DISTRIBUCION: 1.0,
            })

    total_data = {
        COL_CECO2: "TOTAL_GENERAL",
        "CANTIDAD_FILAS": int(summary_df["CANTIDAD_FILAS"].sum()),
        COL_CANTIDAD_EQUIVALENTE: float(summary_df[COL_CANTIDAD_EQUIVALENTE].sum()),
        COL_PORCENTAJE_DISTRIBUCION: 1.0,
    }

    if has_filter:
        total_data[COL_GRUPO_FILTRO] = "TOTAL_GENERAL"

    total_rows.append(total_data)

    summary_df = pd.concat([summary_df, pd.DataFrame(total_rows)], ignore_index=True)
    return summary_df


def build_driver_group_summary(df_driver: pd.DataFrame, driver_type: str) -> pd.DataFrame:
    base_columns = ["DRIVER", "CECOS_UNICOS", "CANTIDAD_FILAS"]

    if driver_type == DRIVER_CANTIDAD:
        base_columns.append("IMPORTE_TOTAL")
    else:
        base_columns.extend([COL_CANTIDAD_EQUIVALENTE, COL_PORCENTAJE_DISTRIBUCION])

    if df_driver is None or df_driver.empty or COL_CECO2 not in df_driver.columns:
        return pd.DataFrame(columns=base_columns)

    work = df_driver.copy()
    work[COL_CECO2] = work[COL_CECO2].apply(clean_cell_text)

    if COL_GRUPO_FILTRO in work.columns:
        work["DRIVER"] = work[COL_GRUPO_FILTRO].apply(clean_cell_text)
    else:
        work["DRIVER"] = "GENERAL"

    if driver_type == DRIVER_CANTIDAD and COL_PRECIO_DISTRIBUIDO in work.columns:
        work["_IMPORTE_TOTAL"] = pd.to_numeric(
            work[COL_PRECIO_DISTRIBUIDO],
            errors="coerce",
        ).fillna(0)

        summary_df = (
            work.groupby("DRIVER", dropna=False)
            .agg(
                CECOS_UNICOS=(COL_CECO2, lambda values: values.dropna().nunique()),
                CANTIDAD_FILAS=(COL_CECO2, "size"),
                IMPORTE_TOTAL=("_IMPORTE_TOTAL", "sum"),
            )
            .reset_index()
        )

        summary_df = summary_df.sort_values("DRIVER").reset_index(drop=True)

        total_row = pd.DataFrame([{
            "DRIVER": "TOTAL_GENERAL",
            "CECOS_UNICOS": int(work[COL_CECO2].dropna().nunique()),
            "CANTIDAD_FILAS": int(summary_df["CANTIDAD_FILAS"].sum()),
            "IMPORTE_TOTAL": float(summary_df["IMPORTE_TOTAL"].sum()),
        }])

        return pd.concat([summary_df, total_row], ignore_index=True)

    work["_CANTIDAD_EQUIVALENTE"] = pd.to_numeric(
        work[COL_PORCENTAJE],
        errors="coerce",
    ).fillna(0)

    summary_df = (
        work.groupby("DRIVER", dropna=False)
        .agg(
            CECOS_UNICOS=(COL_CECO2, lambda values: values.dropna().nunique()),
            CANTIDAD_FILAS=(COL_CECO2, "size"),
            **{COL_CANTIDAD_EQUIVALENTE: ("_CANTIDAD_EQUIVALENTE", "sum")},
        )
        .reset_index()
    )

    total_equivalente = float(summary_df[COL_CANTIDAD_EQUIVALENTE].sum())
    summary_df[COL_PORCENTAJE_DISTRIBUCION] = [
        (float(value) / total_equivalente) if total_equivalente else 0
        for value in summary_df[COL_CANTIDAD_EQUIVALENTE]
    ]

    summary_df = summary_df.sort_values("DRIVER").reset_index(drop=True)

    total_row = pd.DataFrame([{
        "DRIVER": "TOTAL_GENERAL",
        "CECOS_UNICOS": int(work[COL_CECO2].dropna().nunique()),
        "CANTIDAD_FILAS": int(summary_df["CANTIDAD_FILAS"].sum()),
        COL_CANTIDAD_EQUIVALENTE: total_equivalente,
        COL_PORCENTAJE_DISTRIBUCION: 1.0,
    }])

    return pd.concat([summary_df, total_row], ignore_index=True)

def safe_excel_value(value):
    try:
        if pd.isna(value):
            return None
    except Exception:
        pass

    return value


def delete_sheet_if_exists(workbook, sheet_name: str):
    if sheet_name in workbook.sheetnames:
        ws = workbook[sheet_name]
        workbook.remove(ws)


def write_dataframe_to_workbook_sheet(workbook, sheet_name: str, df: pd.DataFrame):
    clean_sheet_name = sanitize_sheet_name(sheet_name)

    delete_sheet_if_exists(workbook, clean_sheet_name)

    ws = workbook.create_sheet(clean_sheet_name)

    header_fill = PatternFill("solid", fgColor="5E6F32")
    header_font = Font(color="FFFFFF", bold=True)

    formats = {
        COL_PORCENTAJE: "0.00%", COL_PORCENTAJE_DISTRIBUCION: "0.00%",
        COL_CANTIDAD_EQUIVALENTE: "#,##0.00", COL_PRECIO: "#,##0.00",
        COL_PRECIO_DISTRIBUIDO: "#,##0.00", "IMPORTE_TOTAL": "#,##0.00",
    }
    column_formats = [formats.get(header) for header in df.columns]
    widths = [max(12, len(str(header)) + 2) for header in df.columns]

    for col_idx, col_name in enumerate(df.columns, start=1):
        cell = ws.cell(row=1, column=col_idx)
        cell.value = col_name
        cell.fill = header_fill
        cell.font = header_font

    for row_idx, row_values in enumerate(df.itertuples(index=False, name=None), start=2):
        for col_idx, value in enumerate(row_values, start=1):
            cell = ws.cell(row=row_idx, column=col_idx)
            cell.value = safe_excel_value(value)
            if column_formats[col_idx - 1]:
                cell.number_format = column_formats[col_idx - 1]
            if row_idx <= 200 and cell.value is not None:
                widths[col_idx - 1] = max(widths[col_idx - 1], len(str(cell.value)) + 2)

    ws.freeze_panes = "A2"

    if ws.max_row >= 1 and ws.max_column >= 1:
        ws.auto_filter.ref = ws.dimensions

    for col_idx, width in enumerate(widths, start=1):
        column_letter = ws.cell(row=1, column=col_idx).column_letter
        ws.column_dimensions[column_letter].width = min(width, 45)


def export_driver_result(
    df_driver: pd.DataFrame,
    df_observados: pd.DataFrame,
    summary: dict,
    output_dir: str | Path,
    output_name: str = "driver_presupuesto",
    source_file_path: str | Path | None = None,
) -> Path:
    df_resumen = pd.DataFrame(
        [{"METRICA": key, "VALOR": value} for key, value in summary.items()]
    )

    driver_type = summary.get("tipo_driver", DRIVER_CANTIDAD)

    df_resumen_ceco = build_ceco_summary(
        df_driver=df_driver,
        driver_type=driver_type,
    )

    df_resumen_driver = build_driver_group_summary(
        df_driver=df_driver,
        driver_type=driver_type,
    )

    # Modo principal: escribir dentro del mismo Excel origen.
    if source_file_path is not None:
        original_path = Path(source_file_path)

        def modify(workbook):
            write_dataframe_to_workbook_sheet(workbook, "DRIVER", df_driver)
            write_dataframe_to_workbook_sheet(workbook, "RESUMEN_DRIVER", df_resumen_driver)
            write_dataframe_to_workbook_sheet(workbook, "RESUMEN_CECO", df_resumen_ceco)
            write_dataframe_to_workbook_sheet(workbook, "OBSERVADOS", df_observados)
            write_dataframe_to_workbook_sheet(workbook, "RESUMEN", df_resumen)

        return atomic_update(df_driver, original_path, modify)

    # Modo legacy: generar archivo nuevo si algun flujo antiguo lo usa.
    output_path = Path(output_dir)
    output_path.mkdir(parents=True, exist_ok=True)

    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    file_path = output_path / f"{output_name}_{timestamp}_{uuid4().hex[:8]}.xlsx"

    def serialize(temporary):
        with pd.ExcelWriter(temporary, engine="xlsxwriter") as writer:
            write_sheets(writer)

    def write_sheets(writer):
        sheets = [
            ("DRIVER", df_driver),
            ("RESUMEN_DRIVER", df_resumen_driver),
            ("RESUMEN_CECO", df_resumen_ceco),
            ("OBSERVADOS", df_observados),
            ("RESUMEN", df_resumen),
        ]

        for sheet_name, df in sheets:
            clean_sheet = sanitize_sheet_name(sheet_name)
            df.to_excel(writer, index=False, sheet_name=clean_sheet)
            apply_xlsxwriter_format(writer, df, clean_sheet)

    return atomic_write(file_path, serialize)
