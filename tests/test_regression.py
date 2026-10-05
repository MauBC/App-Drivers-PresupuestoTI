from copy import deepcopy
from pathlib import Path

import pandas as pd
import pytest
from openpyxl import load_workbook

from conftest import source_config
from src.core.search_executor import run_search, export_result
from src.core.completion_executor import run_completion, export_completion_result
from src.core.driver_executor import run_driver, export_driver_result


def test_exact_match_duplicate_empty_and_sources_unchanged(search_config):
    paths = [Path(search_config[role]["archivo"]["ruta"]) for role in ("template", "base")]
    before = [path.read_bytes() for path in paths]
    df, summary = run_search(search_config)
    assert df["ESTADO_MATCH"].tolist() == ["EXACTO", "NO ENCONTRADO", "NO ENCONTRADO"]
    assert df["Importe"].tolist() == [10, "", ""]
    assert df["VALIDADO"].tolist() == [1, 0, 0]
    assert list(df.columns)[:4] == ["Clave", "Detalle", "Clave_BASE", "Importe"]
    assert summary == {"template_rows": 3, "base_rows": 3, "result_rows": 3,
                       "exactos": 1, "aproximados": 0, "no_encontrados": 2}
    assert [path.read_bytes() for path in paths] == before


@pytest.mark.parametrize("approximate,expected", [(80, "APROXIMADO"), (95, "NO ENCONTRADO")])
def test_similarity_thresholds(excel_file, approximate, expected):
    template = excel_file("consulta.xlsx", ["Nombre"], [["MARIA LOPEZ"], ["LOPEZ MARTA"]])
    base = excel_file("base.xlsx", ["Nombre"], [["MARTA LOPEZ"]])
    config = {"template": source_config(template, ["Nombre"]), "base": source_config(base, ["Nombre"]),
              "match": {"tipo_busqueda": "similaridad", "score_exacto": 100, "score_aproximado": approximate},
              "salida": {"incluir_columnas_tecnicas": True}}
    df, _ = run_search(config)
    assert df["ESTADO_MATCH"].tolist() == [expected, "EXACTO"]
    assert df["SCORE_MATCH"].tolist() == [90.91, 100]


@pytest.mark.parametrize("mode", ["excel_nuevo", "nueva_hoja_template"])
def test_search_exports_preserve_original_sheets(search_config, mode):
    config = deepcopy(search_config)
    config["salida"]["modo"] = mode
    template = Path(config["template"]["archivo"]["ruta"])
    base = Path(config["base"]["archivo"]["ruta"])
    before_template, before_base = template.read_bytes(), base.read_bytes()
    df, _ = run_search(config)
    output = export_result(df, config)
    assert base.read_bytes() == before_base
    if mode == "excel_nuevo":
        assert output != template
        assert template.read_bytes() == before_template
    workbook = load_workbook(output)
    try:
        assert workbook["Cruce"]["D2"].value == 10
        assert workbook["Cruce"].freeze_panes == "A2"
        if mode == "nueva_hoja_template":
            assert workbook["Conservar"]["A1"].value == "=1+2"
            assert workbook["Conservar"].sheet_state == "hidden"
            assert workbook["Datos"]["B3"].value == "uno"
    finally:
        workbook.close()


@pytest.mark.parametrize("mode", ["excel_nuevo", "hoja_nueva_mismo_excel"])
def test_completion_only_reprocesses_zero(search_config, excel_file, mode):
    initial, _ = run_search(search_config)
    initial.loc[1, "VALIDADO"] = 1
    initial.loc[1, "Importe"] = 123
    initial.loc[2, "VALOR_BUSCADO"] = "NUEVO"
    previous = excel_file("previo.xlsx", list(initial.columns), list(initial.itertuples(index=False, name=None)))
    base = excel_file("otra_base.xlsx", ["Clave", "Importe"], [["NUEVO", 45]])
    config = {"previous_result": {"archivo": {"ruta": str(previous)}, "hoja": {"nombre_detectado": "Datos"},
                                  "fila_header": 1, "columna_busqueda": "VALOR_BUSCADO"},
              "base": source_config(base, ["Clave", "Importe"]), "match": search_config["match"],
              "column_mappings": [{"base_alias": "Importe", "base_header_detectado": "Importe", "destination_column_ref": "Importe"}],
              "completion_output": {"modo": mode, "nombre_hoja": "Completado"}}
    before = base.read_bytes()
    result, summary = run_completion(config)
    assert result["Importe"].tolist() == [10, 123, 45]
    assert summary["filas_reprocesadas"] == 1
    assert summary["exactos_nuevos"] == 1
    output = export_completion_result(result, config)
    assert base.read_bytes() == before
    workbook = load_workbook(output)
    try:
        assert workbook["Completado"]["D4"].value == 45
        if mode == "hoja_nueva_mismo_excel":
            assert workbook["Conservar"]["A1"].value == "=1+2"
    finally:
        workbook.close()


@pytest.mark.parametrize("driver_type", ["cantidad", "porcentaje"])
def test_driver_rules_expansion_and_exports(excel_file, driver_type):
    result_path = excel_file("resultado.xlsx", ["DNI", "CECO", "PRECIO", "VALIDADO"], [
        ["001", "ORIG", 100, 1], ["002", "ESP01", 50, 1],
        ["999", "OTRO", 20, 1], ["003", "IGNORAR", 30, 0],
    ])
    base = excel_file("maestro.xlsx", ["DNI", "CECO", "PORCENTAJE"], [["001", "A", .25], ["001", "B", .75]])
    config = {"driver_type": driver_type,
              "resultado": {"file_path": str(result_path), "sheet_name": "Datos", "header_row": 1,
                            "dni_col": "DNI", "ceco_col": "CECO", "price_col": "PRECIO"},
              "base": {"file_path": str(base), "sheet_name": "Datos", "header_row": 1,
                       "dni_col": "DNI", "ceco_col": "CECO", "percentage_col": "PORCENTAJE"},
              "special_ceco_rules": [{"prefijo": "ESP", "ceco2": "ESPECIAL"}]}
    before = base.read_bytes()
    df, observed, summary = run_driver(config)
    assert df["CECO2"].tolist() == ["A", "B", "ESPECIAL", "OTRO"]
    assert df["PORCENTAJE"].tolist() == [.25, .75, 1, 1]
    assert df["FILA_ORIGEN"].tolist() == [2, 2, 3, 4]
    assert len(observed) == 1
    assert summary["filas_no_validadas"] == 1
    if driver_type == "cantidad":
        assert df["PRECIO_DISTRIBUIDO"].tolist() == [25, 75, 50, 20]
    output = export_driver_result(df, observed, summary, result_path.parent, source_file_path=result_path)
    assert base.read_bytes() == before
    workbook = load_workbook(output)
    try:
        assert set(workbook.sheetnames) == {"Datos", "Conservar", "DRIVER", "OBSERVADOS", "RESUMEN", "RESUMEN_DRIVER", "RESUMEN_CECO"}
        assert workbook["Conservar"]["A1"].value == "=1+2"
        assert workbook["DRIVER"]["E2"].number_format == "0.00%"
    finally:
        workbook.close()
