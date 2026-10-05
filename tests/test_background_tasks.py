from pathlib import Path
from threading import Event, get_ident
import time

import customtkinter as ctk
import pytest

from src.gui.background_tasks import BackgroundTasks, tasks_for
from src.gui.search_builder_view import SearchBuilderView
from src.gui import search_builder_view as search_gui
from src.gui.widgets import completion_panel as completion_gui
from src.gui.widgets import driver_panel as driver_gui
from src.core.search_executor import run_search
from conftest import source_config


def pump(root, condition, timeout=20):
    # Verifica comportamiento, no velocidad: margen para Tk/Windows bajo carga.
    deadline = time.monotonic() + timeout
    while not condition() and time.monotonic() < deadline:
        root.update()
        time.sleep(.01)
    assert condition(), "El trabajo no completo la etapa esperada"


@pytest.mark.parametrize("fail", [False, True])
def test_worker_keeps_ui_responsive_and_callbacks_on_main_thread(tk_root, fail):
    manager = BackgroundTasks(tk_root)
    button = ctk.CTkButton(tk_root, text="Ejecutar")
    release, entered = Event(), Event()
    main_thread = get_ident()
    worker_threads, callback_threads, stages, outcomes, heartbeat = [], [], [], [], []
    def work(report):
        worker_threads.append(get_ident())
        report("Procesando...")
        entered.set()
        if not release.wait(5):
            raise RuntimeError("timeout de prueba")
        if fail:
            raise ValueError("error controlado")
        return 42
    def callback(value):
        callback_threads.append(get_ident())
        outcomes.append(value)
    def progress(value):
        callback_threads.append(get_ident())
        stages.append(value)
    try:
        manager.start(work, button, "Ejecutar", callback, callback, progress)
        tk_root.after(10, lambda: heartbeat.append(True))
        pump(tk_root, lambda: entered.is_set() and heartbeat and stages)
        assert manager.busy and button.cget("state") == "disabled"
        with pytest.raises(RuntimeError, match="proceso en curso"):
            manager.start(work, button, "Ejecutar", callback, callback, progress)
        release.set()
        pump(tk_root, lambda: not manager.busy)
        manager.thread.join(timeout=2)
        assert worker_threads == [manager.thread.ident] and worker_threads[0] != main_thread
        assert set(callback_threads) == {main_thread}
        assert stages == ["Procesando..."]
        assert isinstance(outcomes[0], ValueError) if fail else outcomes == [42]
        assert button.cget("state") == "normal" and button.cget("text") == "Ejecutar"
    finally:
        release.set()
        manager.thread.join(timeout=6)
        button.destroy()


@pytest.mark.parametrize("flow", ["search", "completion", "driver"])
def test_actual_panel_pipeline_in_background_with_real_excel(tk_root, search_config, excel_file, monkeypatch, flow):
    view = SearchBuilderView(tk_root)
    messages, failures, heartbeat = [], [], []
    monkeypatch.setattr(search_gui.messagebox, "showinfo", lambda title, text: messages.append(text))
    monkeypatch.setattr(search_gui.messagebox, "showerror", lambda title, text: failures.append(text))
    if flow == "search":
        panel, module, function = view, search_gui, "run_search"
        monkeypatch.setattr(view, "build_runtime_config", lambda: search_config)
        button, execute = view.btn_execute, view.execute_search
        outputs = lambda: list(Path(search_config["template"]["archivo"]["ruta"]).parent.glob("*_resultado_busqueda_*.xlsx"))
    elif flow == "completion":
        initial, _ = run_search(search_config)
        previous = excel_file("previo.xlsx", list(initial.columns), list(initial.itertuples(index=False, name=None)))
        config = {"previous_result": {"archivo": {"ruta": str(previous)}, "hoja": {"nombre_detectado": "Datos"},
                                      "fila_header": 1, "columna_busqueda": "VALOR_BUSCADO"},
                  "base": search_config["base"], "match": search_config["match"],
                  "column_mappings": [{"base_alias": "Importe", "base_header_detectado": "Importe", "destination_column_ref": "Importe"}],
                  "completion_output": {"modo": "excel_nuevo", "nombre_hoja": "Completado"}}
        panel, module, function = view.completion_panel, completion_gui, "run_completion"
        monkeypatch.setattr(panel, "build_config", lambda: config)
        button, execute = panel.btn_complete, panel.execute_completion
        outputs = lambda: list(previous.parent.glob("*_completado_*.xlsx"))
    else:
        destination = excel_file("driver.xlsx", ["DNI", "CECO", "PRECIO"], [["001", "A", 100]])
        base = excel_file("driver_base.xlsx", ["DNI", "CECO", "PORCENTAJE"], [["001", "B", 1]])
        config = {"driver_type": "cantidad",
                  "resultado": {"file_path": str(destination), "sheet_name": "Datos", "header_row": 1,
                                "dni_col": "DNI", "ceco_col": "CECO", "price_col": "PRECIO"},
                  "base": {"file_path": str(base), "sheet_name": "Datos", "header_row": 1,
                           "dni_col": "DNI", "ceco_col": "CECO", "percentage_col": "PORCENTAJE"},
                  "output": {"dir": str(destination.parent), "name": "driver"}}
        panel, module, function = view.driver_panel, driver_gui, "run_driver"
        monkeypatch.setattr(panel, "build_driver_config", lambda: config)
        button, execute = panel.btn_execute, panel.execute_driver
        outputs = lambda: [destination]
    original = getattr(module, function)
    release, entered = Event(), Event()
    worker_threads = []
    def delayed(config, progress=None):
        worker_threads.append(get_ident())
        entered.set()
        if not release.wait(5):
            raise RuntimeError("timeout de prueba")
        return original(config, progress=progress)
    monkeypatch.setattr(module, function, delayed)
    manager = tasks_for(view)
    try:
        execute()
        if flow == "search":
            # Editar la configuracion mientras corre no cambia el trabajo iniciado.
            search_config["salida"]["nombre_hoja_resultado"] = "OtroNombre"
        tk_root.after(10, lambda: heartbeat.append(True))
        pump(tk_root, lambda: entered.is_set() and heartbeat)
        assert manager.busy and button.cget("state") == "disabled"
        release.set()
        pump(tk_root, lambda: not manager.busy)
        manager.thread.join(timeout=2)
        assert messages and not failures
        assert outputs() and all(path.is_file() for path in outputs())
        if flow == "search":
            from openpyxl import load_workbook
            workbook = load_workbook(outputs()[0])
            try:
                assert workbook.sheetnames == ["Cruce"]
            finally:
                workbook.close()
        assert worker_threads[0] != get_ident()
        assert button.cget("state") == "normal"
        assert "Guardando" in view.log_textbox.get("1.0", "end")
    finally:
        release.set()
        if manager.thread:
            manager.thread.join(timeout=6)
        view.close_resources()
        view.destroy()


def test_panel_reports_processing_failure_and_can_retry(tk_root, search_config, monkeypatch):
    view = SearchBuilderView(tk_root)
    errors, results, callback_threads = [], [], []
    monkeypatch.setattr(view, "build_runtime_config", lambda: search_config)
    def show_error(title, text):
        errors.append(text)
        callback_threads.append(get_ident())
    monkeypatch.setattr(search_gui.messagebox, "showerror", show_error)
    monkeypatch.setattr(search_gui.messagebox, "showinfo", lambda title, text: results.append(text))
    source = Path(search_config["base"]["archivo"]["ruta"])
    content = source.read_bytes()
    source.unlink()
    manager = tasks_for(view)
    try:
        view.execute_search()
        pump(tk_root, lambda: not manager.busy)
        assert errors and "No existe" in errors[0]
        assert callback_threads == [get_ident()]
        assert view.btn_execute.cget("state") == "normal"
        source.write_bytes(content)
        view.execute_search()
        pump(tk_root, lambda: not manager.busy)
        assert results
    finally:
        if manager.thread:
            manager.thread.join(timeout=6)
        view.close_resources()
        view.destroy()


def test_panel_reports_actual_suffix_when_requested_sheet_is_input(tk_root, search_config, monkeypatch):
    view = SearchBuilderView(tk_root)
    messages, failures = [], []
    search_config["salida"].update(modo="nueva_hoja_template", nombre_hoja_resultado="Datos")
    monkeypatch.setattr(view, "build_runtime_config", lambda: search_config)
    monkeypatch.setattr(search_gui.messagebox, "showinfo", lambda title, text: messages.append(text))
    monkeypatch.setattr(search_gui.messagebox, "showerror", lambda title, text: failures.append(text))
    manager = tasks_for(view)
    try:
        view.execute_search()
        pump(tk_root, lambda: not manager.busy)
        assert not failures and "Hojas: Datos_2" in messages[0]
        assert "Hojas generadas: Datos_2" in view.log_textbox.get("1.0", "end")
        from openpyxl import load_workbook
        workbook = load_workbook(search_config["template"]["archivo"]["ruta"])
        try:
            assert workbook["Datos"]["A1"].value == "Titulo"
            assert workbook["Datos_2"]["D2"].value == 10
        finally:
            workbook.close()
    finally:
        if manager.thread:
            manager.thread.join(timeout=6)
        view.close_resources()
        view.destroy()


def test_app_cannot_close_while_persisting():
    from app import PresupuestoApp
    class View:
        def __init__(self):
            self.messages, self.closed = [], False
        def log(self, message):
            self.messages.append(message)
        def close_resources(self):
            self.closed = True
    class App:
        def __init__(self):
            self.view = View()
            self._background_tasks = type("Tasks", (), {"busy": True})()
            self.destroyed = False
        def destroy(self):
            self.destroyed = True
    app = App()
    PresupuestoApp.on_close(app)
    assert not app.destroyed and not app.view.closed and app.view.messages
    app._background_tasks.busy = False
    PresupuestoApp.on_close(app)
    assert app.destroyed and app.view.closed
