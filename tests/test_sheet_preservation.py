from pathlib import Path

import pandas as pd
import pytest
from openpyxl import load_workbook
from openpyxl.styles import Font

from src.core.search_executor import run_search, export_result
from src.core.driver_executor import run_driver, export_driver_result
from src.core.completion_executor import export_completion_result
from src.core.excel_safety import captured_sources, attach_capture
from src.core.search_executor import unique_sheet_name


def sheet_signature(ws):
    return {
        "values": list(ws.values), "state": ws.sheet_state,
        "cells": [(c.coordinate, c.value, c.number_format, c.font.bold, c.fill.fgColor.rgb)
                  for row in ws.iter_rows() for c in row],
        "widths": {name: dimension.width for name, dimension in ws.column_dimensions.items()},
    }


@pytest.mark.parametrize("existing,requested,expected", [
    ("Datos", "Datos", "Datos_2"),
    ("Cruce", "Cruce", "Cruce_2"),
    ("Cruce", "cRuCe", "cRuCe_2"),
    ("Cruce_A", "Cruce/A", "Cruce_A_2"),
    ("A" * 31, "A" * 40, "A" * 29 + "_2"),
])
def test_search_preserves_existing_sheets_and_reports_free_name(search_config, existing, requested, expected):
    path = Path(search_config["template"]["archivo"]["ruta"])
    workbook = load_workbook(path)
    try:
        if existing != "Datos":
            old = workbook.create_sheet(existing)
            old["A1"] = "Datos previos que no deben borrarse"
            old["B2"] = "=1+2"
            old["B2"].font = Font(bold=True)
            old.column_dimensions["B"].width = 34
            old.sheet_state = "hidden"
        signatures = {name: sheet_signature(workbook[name]) for name in workbook.sheetnames}
        workbook.save(path)
    finally:
        workbook.close()
    search_config["salida"].update(modo="nueva_hoja_template", nombre_hoja_resultado=requested)
    df, _ = run_search(search_config)
    export_result(df, search_config)
    workbook = load_workbook(path)
    try:
        for name, signature in signatures.items():
            assert sheet_signature(workbook[name]) == signature
        assert expected in workbook.sheetnames
        assert workbook[expected]["D2"].value == 10
        assert len(expected) <= 31
    finally:
        workbook.close()
    assert df.attrs["excel_output_sheets"] == [expected]


def test_repeated_search_keeps_previous_results(search_config):
    path = Path(search_config["template"]["archivo"]["ruta"])
    search_config["salida"]["modo"] = "nueva_hoja_template"
    first, _ = run_search(search_config)
    export_result(first, search_config)
    workbook = load_workbook(path)
    try:
        original = sheet_signature(workbook["Cruce"])
    finally:
        workbook.close()
    second, _ = run_search(search_config)
    export_result(second, search_config)
    workbook = load_workbook(path)
    try:
        assert sheet_signature(workbook["Cruce"]) == original
        assert "Cruce_2" in workbook.sheetnames
    finally:
        workbook.close()
    assert second.attrs["excel_output_sheets"] == ["Cruce_2"]


@pytest.mark.parametrize("source_sheet", ["DRIVER", "RESUMEN", "Datos"])
def test_driver_preserves_source_and_all_preexisting_output_sheets(excel_file, source_sheet):
    path = excel_file("resultado.xlsx", ["DNI", "CECO", "PRECIO"], [["001", "ORIG", 100]])
    base = excel_file("maestro.xlsx", ["DNI", "CECO", "PORCENTAJE"], [["001", "NUEVO", .5]])
    outputs = ["DRIVER", "RESUMEN_DRIVER", "RESUMEN_CECO", "OBSERVADOS", "RESUMEN"]
    workbook = load_workbook(path)
    try:
        workbook["Datos"].title = source_sheet
        for name in outputs:
            if name != source_sheet:
                sheet = workbook.create_sheet(name)
                sheet["A1"] = "Conservar " + name
        original = {name: sheet_signature(workbook[name]) for name in workbook.sheetnames}
        workbook.save(path)
    finally:
        workbook.close()
    config = {"driver_type": "cantidad",
              "resultado": {"file_path": str(path), "sheet_name": source_sheet, "header_row": 1,
                            "dni_col": "DNI", "ceco_col": "CECO", "price_col": "PRECIO"},
              "base": {"file_path": str(base), "sheet_name": "Datos", "header_row": 1,
                       "dni_col": "DNI", "ceco_col": "CECO", "percentage_col": "PORCENTAJE"}}
    df, observed, summary = run_driver(config)
    export_driver_result(df, observed, summary, path.parent, source_file_path=path)
    workbook = load_workbook(path)
    try:
        for name, signature in original.items():
            assert sheet_signature(workbook[name]) == signature
        assert all(name + "_2" in workbook.sheetnames for name in outputs)
        assert workbook["DRIVER_2"]["D2"].value == "NUEVO"
    finally:
        workbook.close()
    assert df.attrs["excel_output_sheets"] == [name + "_2" for name in outputs]


def test_completion_case_insensitive_collision_is_explicit(excel_file):
    path = excel_file("resultado.xlsx", ["Columna"], [["Original"]])
    df = pd.DataFrame({"Columna": ["Resultado nuevo"]})
    with captured_sources({"previous_result": path}) as sources:
        attach_capture(df, sources, "previous_result")
    config = {"previous_result": {"archivo": {"ruta": str(path)}},
              "completion_output": {"modo": "hoja_nueva_mismo_excel", "nombre_hoja": "dAtOs"}}
    export_completion_result(df, config)
    workbook = load_workbook(path)
    try:
        assert workbook["Datos"]["A2"].value == "Original"
        assert workbook["dAtOs_2"]["A2"].value == "Resultado nuevo"
    finally:
        workbook.close()
    assert df.attrs["excel_output_sheets"] == ["dAtOs_2"]


def test_long_name_skips_existing_suffixes_and_keeps_excel_limit():
    from openpyxl import Workbook
    workbook = Workbook()
    try:
        base = "A" * 31
        workbook.active.title = base
        for number in range(2, 10):
            workbook.create_sheet("A" * 29 + f"_{number}")
        workbook.create_sheet("A" * 28 + "_10")
        assert unique_sheet_name(workbook, base) == "A" * 28 + "_11"
    finally:
        workbook.close()
