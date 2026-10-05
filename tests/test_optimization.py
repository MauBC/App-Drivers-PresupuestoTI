from copy import deepcopy
from pathlib import Path

import pandas as pd
import pytest
from openpyxl import Workbook, load_workbook

from conftest import source_config
from src.core import search_executor as search
from src.core import completion_executor as completion
from src.core.driver_executor import write_dataframe_to_workbook_sheet


def test_repeated_similarity_keeps_rows_scores_and_first_duplicate(excel_file, monkeypatch):
    template = excel_file("consultas.xlsx", ["Nombre", "Detalle"], [
        ["María López", "primero"], [" MARIA LOPEZ ", "segundo"],
        ["MARIA LOPEZ", "tercero"], [None, "vacio"],
    ])
    base = excel_file("maestro.xlsx", ["Nombre", "Valor"], [["MARTA LOPEZ", 10], ["MARTA LOPEZ", 99]])
    config = {"template": source_config(template, ["Nombre", "Detalle"]),
              "base": source_config(base, ["Nombre", "Valor"]),
              "match": {"tipo_busqueda": "similaridad", "score_exacto": 100, "score_aproximado": 80},
              "salida": {"incluir_columnas_tecnicas": True}}
    original = search.find_best_similarity
    calls = []
    def counted(query, choices):
        calls.append(query)
        return original(query, choices)
    monkeypatch.setattr(search, "find_best_similarity", counted)
    df, summary = search.run_search(config)
    assert calls == ["MARIA LOPEZ"]
    assert df["ESTADO_MATCH"].tolist() == ["APROXIMADO"] * 3 + ["NO ENCONTRADO"]
    assert df["SCORE_MATCH"].tolist() == [90.91] * 3 + [0]
    assert df["Valor"].tolist() == [10, 10, 10, ""]
    assert df["Detalle"].tolist() == ["primero", "segundo", "tercero", "vacio"]
    assert summary["result_rows"] == 4
    calls.clear()
    workbook = load_workbook(base)
    try:
        workbook["Datos"]["B2"] = 25
        workbook.save(base)
    finally:
        workbook.close()
    changed, _ = search.run_search(config)
    assert changed["Valor"].tolist()[:3] == [25, 25, 25]
    assert calls == ["MARIA LOPEZ"]  # Cache no persiste entre ejecuciones.


def test_completion_caches_only_matching_not_mapped_rows(search_config, excel_file, monkeypatch):
    initial, _ = search.run_search(search_config)
    initial["VALIDADO"] = 0
    initial["VALOR_BUSCADO"] = ["María López", " MARIA LOPEZ ", "MARIA LOPEZ"]
    previous = excel_file("previo.xlsx", list(initial.columns), list(initial.itertuples(index=False, name=None)))
    base = excel_file("maestro2.xlsx", ["Clave", "Importe"], [["MARTA LOPEZ", 45]])
    config = {"previous_result": {"archivo": {"ruta": str(previous)}, "hoja": {"nombre_detectado": "Datos"},
                                  "fila_header": 1, "columna_busqueda": "VALOR_BUSCADO"},
              "base": source_config(base, ["Clave", "Importe"]),
              "match": {"tipo_busqueda": "similaridad", "score_exacto": 100, "score_aproximado": 80},
              "column_mappings": [{"base_alias": "Importe", "base_header_detectado": "Importe", "destination_column_ref": "Importe"}]}
    original = completion.find_best_similarity
    calls = []
    def counted(query, choices):
        calls.append(query)
        return original(query, choices)
    monkeypatch.setattr(completion, "find_best_similarity", counted)
    df, summary = completion.run_completion(config)
    assert calls == ["MARIA LOPEZ"]
    assert df["Importe"].tolist() == [45, 45, 45]
    assert df["VALIDADO"].tolist() == [1, 1, 1]
    assert summary["aproximados_nuevos"] == 3


def test_driver_formatting_width_sampling_and_empty_values():
    df = pd.DataFrame({"TEXTO": ["corto"] * 199 + ["x" * 100],
                       "PORCENTAJE": [.25] * 200, "PRECIO": [10] * 199 + [None],
                       "CANTIDAD_EQUIVALENTE": [1.] * 200})
    workbook = Workbook()
    try:
        write_dataframe_to_workbook_sheet(workbook, "DRIVER", df)
        ws = workbook["DRIVER"]
        assert ws.column_dimensions["A"].width == 12
        assert ws["A201"].value == "x" * 100  # Fuera de muestra de ancho.
        assert ws["B201"].number_format == "0.00%"
        assert ws["C201"].value is None and ws["C201"].number_format == "#,##0.00"
        assert ws["D2"].number_format == "#,##0.00"
        assert ws["A1"].font.bold
        assert ws.freeze_panes == "A2" and ws.auto_filter.ref == "A1:D201"
    finally:
        workbook.close()


def test_progress_stages_do_not_change_result(search_config):
    stages = []
    df, summary = search.run_search(search_config, progress=stages.append)
    reference, expected = search.run_search(deepcopy(search_config))
    pd.testing.assert_frame_equal(df, reference)
    assert summary == expected
    assert stages == ["Capturando template y maestro...", "Validando encabezados...", "Procesando busqueda..."]
