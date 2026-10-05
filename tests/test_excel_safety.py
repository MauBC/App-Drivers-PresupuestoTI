from copy import deepcopy
import os
from pathlib import Path
import subprocess
import sys
import time
from zipfile import ZipFile, ZIP_DEFLATED

import pytest
from openpyxl import load_workbook
from openpyxl.styles import Font
from openpyxl.worksheet.table import Table
from openpyxl.workbook.defined_name import DefinedName

from src.core import excel_safety as safety
from src.core.excel_inspector import get_sheet_names, read_headers, read_sheet_sample
from src.core.search_executor import run_search, export_result


def save_copy(source, target):
    workbook = load_workbook(source)
    try:
        workbook["Datos"]["B2"] = "resultado"
        workbook.save(target)
    finally:
        workbook.close()


def change_cell(path, value="cambio humano"):
    workbook = load_workbook(path)
    try:
        workbook["Datos"]["B2"] = value
        workbook.save(path)
    finally:
        workbook.close()


def test_consistent_pair_metadata_and_cleanup(excel_file):
    first = excel_file("a.xlsx", ["Clave"], [[1]])
    second = excel_file("b.xlsx", ["Clave"], [[2]])
    before = [path.read_bytes() for path in (first, second)]
    with safety.captured_sources({"template": first, "base": second}) as snapshots:
        copies = [Path(s.snapshot) for s in snapshots.values()]
        for snapshot in snapshots.values():
            assert snapshot.sha256 == safety.file_hash(snapshot.original) == safety.file_hash(snapshot.snapshot)
            assert snapshot.size > 0 and snapshot.mtime_ns > 0
            assert snapshot.attempt == 1 and snapshot.captured_at.endswith("+00:00")
    assert all(not path.exists() for path in copies)
    assert [path.read_bytes() for path in (first, second)] == before


def test_pair_retries_if_base_changes_during_template_copy(excel_file, monkeypatch):
    first = excel_file("a.xlsx", ["Clave"], [[1]])
    second = excel_file("b.xlsx", ["Clave"], [[2]])
    original_copy = safety.shutil.copyfile
    calls = []
    def copy(source, target):
        calls.append(source)
        result = original_copy(source, target)
        if len(calls) == 1:
            change_cell(second)
        return result
    monkeypatch.setattr(safety.shutil, "copyfile", copy)
    with safety.captured_sources({"template": first, "base": second}, delay=0) as snapshots:
        assert snapshots["template"].attempt == snapshots["base"].attempt == 2
        assert snapshots["base"].sha256 == safety.file_hash(second)
    assert len(calls) == 4


def test_unstable_capture_stops_and_cleans(excel_file, monkeypatch):
    source = excel_file("a.xlsx", ["Clave"], [[1]])
    original_copy = safety.shutil.copyfile
    targets = []
    def copy(path, target):
        original_copy(path, target)
        targets.append(Path(target))
        change_cell(path, len(targets))
    monkeypatch.setattr(safety.shutil, "copyfile", copy)
    with pytest.raises(safety.ExcelChangedError):
        with safety.captured_sources({"template": source}, attempts=2, delay=0):
            pytest.fail("Una captura inestable no debe llegar al match")
    assert len(targets) == 2 and all(not p.exists() for p in targets)


@pytest.mark.parametrize("kind", ["missing", "invalid", "office"])
def test_invalid_sources_rejected(tmp_path, kind):
    source = tmp_path / ("~$datos.xlsx" if kind == "office" else "datos.xlsx")
    if kind != "missing":
        source.write_text("no es Excel")
    with pytest.raises(FileNotFoundError if kind == "missing" else safety.ExcelInvalidError):
        with safety.captured_sources({"template": source}):
            pytest.fail("Entrada invalida")


def test_original_change_after_capture_does_not_change_snapshot(excel_file):
    source = excel_file("a.xlsx", ["Clave", "Valor"], [[1, "antes"]])
    with safety.captured_sources({"template": source}) as snapshots:
        change_cell(source)
        workbook = load_workbook(snapshots["template"].snapshot)
        try:
            assert workbook["Datos"]["B2"].value == "antes"
        finally:
            workbook.close()


@pytest.mark.parametrize("failure", ["serialize", "invalid", "replace", "changed"])
def test_atomic_failure_preserves_original_and_removes_temp(excel_file, monkeypatch, failure):
    destination = excel_file("destino.xlsx", ["Clave", "Valor"], [[1, "original"]])
    before = destination.read_bytes()
    expected = safety.file_hash(destination)
    external = []
    def serialize(temporary):
        if failure == "serialize":
            temporary.write_bytes(b"parcial")
            raise RuntimeError("fallo al guardar")
        if failure == "invalid":
            temporary.write_bytes(b"ZIP incompleto")
            return
        save_copy(destination, temporary)
        if failure == "changed":
            change_cell(destination)
            external.append(destination.read_bytes())
    if failure == "replace":
        def denied(*args):
            raise PermissionError("bloqueado")
        monkeypatch.setattr(safety.os, "replace", denied)
        monkeypatch.setattr(safety.time, "sleep", lambda _: None)
    expected_error = {"serialize": RuntimeError, "invalid": safety.ExcelInvalidError,
                      "replace": safety.ExcelBusyError, "changed": safety.ExcelChangedError}[failure]
    with pytest.raises(expected_error):
        safety.atomic_write(destination, serialize, expected)
    assert destination.read_bytes() == (external[0] if external else before)
    assert not list(destination.parent.glob(".presupuesto_*"))


def test_atomic_success_uses_sibling_and_replace(excel_file, monkeypatch):
    destination = excel_file("destino.xlsx", ["Clave", "Valor"], [[1, "original"]])
    original_replace = safety.os.replace
    replaced = []
    def replace(source, target):
        assert Path(source).parent == Path(target).parent
        safety.validate_package(source)
        replaced.append((source, target))
        return original_replace(source, target)
    monkeypatch.setattr(safety.os, "replace", replace)
    safety.atomic_write(destination, lambda target: save_copy(destination, target), safety.file_hash(destination))
    assert len(replaced) == 1
    assert not list(destination.parent.glob(".presupuesto_*"))
    workbook = load_workbook(destination)
    try:
        assert workbook["Datos"]["B2"].value == "resultado"
    finally:
        workbook.close()


def test_new_destination_does_not_overwrite_file_created_during_save(excel_file, tmp_path):
    source = excel_file("source.xlsx", ["Clave"], [[1]])
    destination = tmp_path / "nuevo.xlsx"
    def serialize(target):
        target.write_bytes(source.read_bytes())
        destination.write_bytes(b"archivo ajeno")
    with pytest.raises(safety.ExcelChangedError):
        safety.atomic_write(destination, serialize)
    assert destination.read_bytes() == b"archivo ajeno"
    assert not list(tmp_path.glob(".presupuesto_*"))


@pytest.mark.parametrize("mutation", ["change", "delete"])
def test_search_refuses_to_overwrite_changed_or_deleted_template(search_config, mutation):
    search_config["salida"]["modo"] = "nueva_hoja_template"
    df, _ = run_search(search_config)
    destination = Path(search_config["template"]["archivo"]["ruta"])
    if mutation == "change":
        change_cell(destination)
        before = destination.read_bytes()
    else:
        destination.unlink()
        before = None
    with pytest.raises(safety.ExcelChangedError):
        export_result(df, search_config)
    assert (destination.read_bytes() if destination.exists() else None) == before


def test_same_master_and_template_cannot_be_updated(search_config):
    search_config["base"] = deepcopy(search_config["template"])
    search_config["salida"]["modo"] = "nueva_hoja_template"
    df, _ = run_search(search_config)
    with pytest.raises(safety.ExcelError, match="maestro"):
        export_result(df, search_config)


def test_headers_changed_after_configuration(search_config):
    path = Path(search_config["template"]["archivo"]["ruta"])
    workbook = load_workbook(path)
    try:
        workbook["Datos"]["A2"] = "OtraColumna"
        workbook.save(path)
    finally:
        workbook.close()
    with pytest.raises(safety.ExcelHeadersChangedError):
        run_search(search_config)


def test_data_can_change_after_configuration(search_config):
    path = Path(search_config["template"]["archivo"]["ruta"])
    workbook = load_workbook(path)
    try:
        workbook["Datos"]["A3"] = "Ausente"
        workbook.save(path)
    finally:
        workbook.close()
    df, summary = run_search(search_config)
    assert summary["exactos"] == 0
    assert df.iloc[0]["VALOR_BUSCADO"] == "Ausente"


def test_office_marker_prevents_writing(search_config, monkeypatch):
    search_config["salida"]["modo"] = "nueva_hoja_template"
    df, _ = run_search(search_config)
    path = Path(search_config["template"]["archivo"]["ruta"])
    before = path.read_bytes()
    path.with_name("~$" + path.name).touch()
    monkeypatch.setattr(safety.time, "sleep", lambda _: None)
    with pytest.raises(safety.ExcelBusyError):
        export_result(df, search_config)
    assert path.read_bytes() == before


@pytest.mark.parametrize("error", [PermissionError("permiso"), OSError("sharing violation")])
def test_transient_retry_bounded_and_backoff(monkeypatch, error):
    if not isinstance(error, PermissionError):
        error.winerror = 32
    attempts, sleeps = [], []
    def fail():
        attempts.append(1)
        raise error
    monkeypatch.setattr(safety.time, "sleep", sleeps.append)
    with pytest.raises(safety.ExcelBusyError):
        safety.retry(fail, "archivo.xlsx", delay=.1)
    assert len(attempts) == 3 and sleeps == [.1, .2]


def test_nontransient_error_not_retried():
    attempts = []
    def fail():
        attempts.append(1)
        raise FileNotFoundError("archivo inexistente")
    with pytest.raises(FileNotFoundError):
        safety.retry(fail, "archivo.xlsx")
    assert len(attempts) == 1


def test_permission_error_before_snapshot_copy_is_reported(excel_file, monkeypatch):
    source = excel_file("archivo.xlsx", ["Clave"], [[1]])
    calls = []
    def denied(path):
        calls.append(path)
        raise PermissionError(13, "ocupado", str(path))
    monkeypatch.setattr(safety, "file_hash", denied)
    with pytest.raises(safety.ExcelBusyError) as caught:
        with safety.captured_sources({"base": source}, delay=0):
            pytest.fail("No se puede procesar una fuente no legible")
    assert len(calls) == 3
    assert caught.value.path == source


@pytest.mark.parametrize("operation", [get_sheet_names,
    lambda p: read_headers(p, "Datos", 1), lambda p: read_sheet_sample(p, "Datos"),
    lambda p: read_headers(p, "NoExiste", 1)])
def test_inspection_closes_handles_even_on_error(excel_file, operation):
    source = excel_file("archivo.xlsx", ["Clave"], [[1]])
    before = source.read_bytes()
    try:
        operation(str(source))
    except ValueError:
        pass
    safety.check_writable(source)  # En Windows exige abrir sin compartir handles.
    renamed = source.with_name("renombrado.xlsx")
    source.rename(renamed)
    assert renamed.read_bytes() == before


def test_source_handles_released_after_match(search_config):
    run_search(search_config)
    for role in ("template", "base"):
        safety.check_writable(search_config[role]["archivo"]["ruta"])


def test_match_uses_snapshots_and_original_config_is_preserved(search_config, monkeypatch):
    from src.core import search_executor
    config_before = deepcopy(search_config)
    original_run = search_executor._run_search
    def run(local):
        assert local["base"]["archivo"]["ruta"] != search_config["base"]["archivo"]["ruta"]
        path = Path(search_config["base"]["archivo"]["ruta"])
        workbook = load_workbook(path)
        try:
            workbook["Datos"]["B2"] = 999
            workbook.save(path)
        finally:
            workbook.close()
        return original_run(local)
    monkeypatch.setattr(search_executor, "_run_search", run)
    df, _ = search_executor.run_search(search_config)
    assert df.iloc[0]["Importe"] == 10
    assert search_config == config_before
    assert df.iloc[0]["FUENTE_ARCHIVO"] == "maestro.xlsx"
    assert all(not Path(source["snapshot"]).exists() for source in df.attrs[safety.CAPTURE_ATTRIBUTE]["sources"].values())


def test_failed_workbook_modification_closes_destination(search_config):
    df, _ = run_search(search_config)
    destination = Path(search_config["template"]["archivo"]["ruta"])
    before = destination.read_bytes()
    def modify(workbook):
        workbook.create_sheet("Parcial")
        raise RuntimeError("fallo en modificacion")
    with pytest.raises(RuntimeError):
        safety.atomic_update(df, destination, modify)
    safety.check_writable(destination)
    assert destination.read_bytes() == before


def test_second_export_with_old_capture_is_rejected(search_config):
    search_config["salida"]["modo"] = "nueva_hoja_template"
    df, _ = run_search(search_config)
    destination = export_result(df, search_config)
    before = destination.read_bytes()
    with pytest.raises(safety.ExcelChangedError):
        export_result(df, search_config)
    assert destination.read_bytes() == before


def test_formats_tables_names_and_hidden_sheet_preserved(search_config):
    path = Path(search_config["template"]["archivo"]["ruta"])
    workbook = load_workbook(path)
    try:
        workbook["Datos"]["B3"].font = Font(bold=True)
        workbook["Datos"].add_table(Table(displayName="Original", ref="A2:B5"))
        workbook.defined_names.add(DefinedName("Referencia", attr_text="'Datos'!$B$3"))
        workbook.save(path)
    finally:
        workbook.close()
    search_config["salida"]["modo"] = "nueva_hoja_template"
    df, _ = run_search(search_config)
    export_result(df, search_config)
    workbook = load_workbook(path)
    try:
        assert workbook["Datos"]["B3"].font.bold
        assert workbook["Datos"].tables["Original"].ref == "A2:B5"
        assert workbook.defined_names["Referencia"].attr_text == "'Datos'!$B$3"
        assert workbook["Conservar"].sheet_state == "hidden"
        assert workbook["Conservar"]["A1"].value == "=1+2"
    finally:
        workbook.close()


def test_vba_package_survives_update(search_config):
    # Fixture de paquete VBA sintetico: comprueba bytes, no ejecucion de macros.
    source = Path(search_config["template"]["archivo"]["ruta"])
    macro_path = source.with_suffix(".xlsm")
    payload = b"fixture-vba-preservar"
    with ZipFile(source) as original, ZipFile(macro_path, "w", ZIP_DEFLATED) as target:
        for item in original.infolist():
            data = original.read(item.filename)
            if item.filename == "[Content_Types].xml":
                data = data.replace(b"application/vnd.openxmlformats-officedocument.spreadsheetml.sheet.main+xml",
                                    b"application/vnd.ms-excel.sheet.macroEnabled.main+xml")
                data = data.replace(b"</Types>", b'<Override PartName="/xl/vbaProject.bin" ContentType="application/vnd.ms-office.vbaProject"/></Types>')
            if item.filename == "xl/_rels/workbook.xml.rels":
                data = data.replace(b"</Relationships>", b'<Relationship Type="http://schemas.microsoft.com/office/2006/relationships/vbaProject" Target="vbaProject.bin" Id="rIdVba"/></Relationships>')
            target.writestr(item.filename, data)
        target.writestr("xl/vbaProject.bin", payload)
    search_config["template"]["archivo"]["ruta"] = str(macro_path)
    search_config["salida"]["modo"] = "nueva_hoja_template"
    df, _ = run_search(search_config)
    export_result(df, search_config)
    with ZipFile(macro_path) as archive:
        assert archive.read("xl/vbaProject.bin") == payload


def test_other_process_cannot_acquire_destination_guard(excel_file):
    source = excel_file("archivo.xlsx", ["Clave"], [[1]])
    script = "from src.core.excel_safety import destination_guard, ExcelBusyError\nimport sys\ntry:\n with destination_guard(sys.argv[1]): pass\nexcept ExcelBusyError:\n sys.exit(17)\n"
    with safety.destination_guard(source):
        result = subprocess.run([sys.executable, "-c", script, str(source)], capture_output=True, timeout=20)
    assert result.returncode == 17, result.stderr.decode()
    with safety.destination_guard(source):
        pass


@pytest.mark.skipif(os.name != "nt", reason="Verificacion real de sharing violation Windows")
def test_windows_open_handle_detected_without_office_marker(excel_file):
    import ctypes
    from ctypes import wintypes
    source = excel_file("archivo.xlsx", ["Clave"], [[1]])
    kernel = ctypes.WinDLL("kernel32", use_last_error=True)
    create = kernel.CreateFileW
    create.argtypes = [wintypes.LPCWSTR, wintypes.DWORD, wintypes.DWORD, ctypes.c_void_p, wintypes.DWORD, wintypes.DWORD, wintypes.HANDLE]
    create.restype = wintypes.HANDLE
    close = kernel.CloseHandle
    close.argtypes = [wintypes.HANDLE]
    handle = create(str(source), 0x80000000, 0, None, 3, 0x80, None)
    assert handle != ctypes.c_void_p(-1).value
    try:
        with pytest.raises(OSError) as caught:
            safety.check_writable(source)
        assert caught.value.winerror == 32
    finally:
        close(handle)
    safety.check_writable(source)


def test_crash_during_serialization_preserves_original_and_releases_guard(excel_file, tmp_path):
    source = excel_file("archivo.xlsx", ["Clave"], [[1]])
    before = source.read_bytes()
    marker = tmp_path / "ready.txt"
    script = """import sys, time
from pathlib import Path
from src.core.excel_safety import atomic_write, file_hash
def serialize(target):
    target.write_bytes(b'parcial')
    Path(sys.argv[2]).write_text(str(target))
    time.sleep(30)
atomic_write(sys.argv[1], serialize, file_hash(sys.argv[1]))
"""
    process = subprocess.Popen([sys.executable, "-c", script, str(source), str(marker)],
                               stdout=subprocess.PIPE, stderr=subprocess.PIPE)
    try:
        deadline = time.monotonic() + 15
        while not marker.exists() and process.poll() is None and time.monotonic() < deadline:
            time.sleep(.05)
        assert marker.exists(), "El proceso no llego al guardado"
        process.kill()
        process.communicate(timeout=10)
        assert source.read_bytes() == before
        with safety.destination_guard(source):
            pass
        # Un kill no ejecuta finally: el temporal huerfano no se publica.
        orphan = Path(marker.read_text())
        assert orphan.exists() and orphan.read_bytes() == b"parcial"
        orphan.unlink()
    finally:
        if process.poll() is None:
            process.kill()
        process.communicate(timeout=10)
