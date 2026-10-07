from pathlib import Path

import pytest
from openpyxl import load_workbook

from src.core import excel_safety as safety, excel_recovery as recovery


def serialize_copy(source, temporary):
    book = load_workbook(source)
    try:
        book.create_sheet("RESULTADO")["A1"] = "Perú / Año: 100"
        book.save(temporary)
    finally:
        book.close()


@pytest.mark.parametrize("code", [5, 32, 33])
@pytest.mark.parametrize("persistent", [False, True])
def test_replace_busy_preserves_result_and_recovers_without_recalculation(excel_file, monkeypatch, code, persistent):
    destination = excel_file("OneDrive.xlsx", ["Clave"], [[1]])
    before = destination.read_bytes()
    expected = safety.file_hash(destination)
    real_replace = safety.os.replace
    attempts, sleeps, calculated = [], [], []
    def replace(source, target):
        attempts.append(target)
        if persistent or len(attempts) < 3:
            error = PermissionError("bloqueo simulado")
            error.winerror = code
            raise error
        real_replace(source, target)
    def serialize(temporary):
        calculated.append(1)
        serialize_copy(destination, temporary)
    monkeypatch.setattr(safety.os, "replace", replace)
    monkeypatch.setattr(safety.time, "sleep", sleeps.append)
    if persistent:
        with pytest.raises(safety.ExcelBusyError, match="Resultados pendientes") as caught:
            safety.atomic_write(destination, serialize, expected)
        assert destination.read_bytes() == before
        assert sleeps == [.5, 1, 1.5, 2]
        assert len(attempts) == 5
        pending = recovery.pending()
        assert len(pending) == 1
        assert pending[0]["id"] == caught.value.recovery_id
        monkeypatch.setattr(safety.os, "replace", real_replace)
        assert recovery.restore(pending[0]["id"]) == destination
        # Otra llamada no vuelve a publicar ni crea otra hoja.
        assert recovery.restore(pending[0]["id"]) == destination
    else:
        safety.atomic_write(destination, serialize, expected)
        assert sleeps == [.5, 1]
    assert calculated == [1]
    assert recovery.pending() == []
    book = load_workbook(destination)
    try:
        assert book.sheetnames.count("RESULTADO") == 1
        assert book["RESULTADO"]["A1"].value == "Perú / Año: 100"
    finally:
        book.close()


def make_pending(excel_file, monkeypatch):
    destination = excel_file("destino.xlsx", ["Clave"], [[1]])
    monkeypatch.setattr(safety, "check_writable", lambda _: (_ for _ in ()).throw(safety.ExcelBusyError(destination)))
    monkeypatch.setattr(safety.time, "sleep", lambda _: None)
    with pytest.raises(safety.ExcelBusyError):
        safety.atomic_write(destination, lambda t: serialize_copy(destination, t), safety.file_hash(destination))
    return destination, recovery.pending()[0]


def test_restart_conflict_and_alternate_preserve_human_edits(excel_file, monkeypatch, tmp_path):
    destination, record = make_pending(excel_file, monkeypatch)
    # El cálculo y su DataFrame ya no existen; sólo se lee el registro del disco.
    monkeypatch.undo()
    monkeypatch.setenv("LOCALAPPDATA", str(tmp_path / "localappdata"))
    book = load_workbook(destination)
    book["Datos"]["A2"] = "edición humana"
    book.save(destination)
    book.close()
    before = destination.read_bytes()
    with pytest.raises(safety.ExcelChangedError):
        recovery.restore(record["id"])
    assert destination.read_bytes() == before
    alternate = tmp_path / "copia.xlsx"
    assert recovery.restore(record["id"], alternate) == alternate
    assert destination.read_bytes() == before
    assert recovery.pending() == []


def test_recovery_refuses_existing_alternate_and_corrupted_payload(excel_file, monkeypatch, tmp_path):
    destination, record = make_pending(excel_file, monkeypatch)
    monkeypatch.undo()
    monkeypatch.setenv("LOCALAPPDATA", str(tmp_path / "localappdata"))
    alternate = excel_file("otra.xlsx", ["Otro"], [[2]])
    before = alternate.read_bytes()
    with pytest.raises(safety.ExcelChangedError):
        recovery.restore(record["id"], alternate)
    assert alternate.read_bytes() == before
    payload = recovery.recovery_directory() / record["id"] / record["payload"]
    payload.write_bytes(b"corrupto")
    with pytest.raises(ValueError, match="dañada"):
        recovery.restore(record["id"])


def test_crash_after_publish_is_reconciled_without_repeating(excel_file, monkeypatch):
    destination = excel_file("destino.xlsx", ["Clave"], [[1]])
    real_complete = recovery.complete
    monkeypatch.setattr(recovery, "complete", lambda *args: (_ for _ in ()).throw(RuntimeError("crash antes del recibo")))
    with pytest.raises(RuntimeError):
        safety.atomic_write(destination, lambda t: serialize_copy(destination, t), safety.file_hash(destination))
    record = recovery.pending()[0]
    before = destination.read_bytes()
    monkeypatch.setattr(recovery, "complete", real_complete)
    assert recovery.restore(record["id"]) == destination
    assert destination.read_bytes() == before
    assert recovery.pending() == []


def test_pending_does_not_block_other_jobs(excel_file, monkeypatch, tmp_path):
    _, record = make_pending(excel_file, monkeypatch)
    monkeypatch.undo()
    monkeypatch.setenv("LOCALAPPDATA", str(tmp_path / "localappdata"))
    source = excel_file("libre.xlsx", ["Clave"], [[3]])
    safety.atomic_write(tmp_path / "nuevo.xlsx", lambda t: serialize_copy(source, t))
    assert [r["id"] for r in recovery.pending()] == [record["id"]]


def test_folder_permissions_keep_local_copy_for_alternate(excel_file, monkeypatch, tmp_path):
    source = excel_file("fuente.xlsx", ["Clave"], [[1]])
    destination = tmp_path / "sin_permiso" / "salida.xlsx"
    real_mkstemp = safety.tempfile.mkstemp
    def mkstemp(*args, **kwargs):
        if kwargs.get("dir") == destination.parent:
            raise PermissionError("carpeta de destino sin permiso")
        return real_mkstemp(*args, **kwargs)
    monkeypatch.setattr(safety.tempfile, "mkstemp", mkstemp)
    monkeypatch.setattr(safety.time, "sleep", lambda _: None)
    with pytest.raises(safety.ExcelBusyError):
        safety.atomic_write(destination, lambda t: serialize_copy(source, t))
    record = recovery.pending()[0]
    assert not destination.exists()
    alternate = tmp_path / "recuperado.xlsx"
    assert recovery.restore(record["id"], alternate) == alternate


def test_recovery_panel_can_save_after_restart(tk_root, excel_file, monkeypatch, tmp_path):
    from src.gui.recovery_panel import RecoveryPanel
    from src.gui.background_tasks import tasks_for
    from src.gui import recovery_panel as gui
    import time
    destination, record = make_pending(excel_file, monkeypatch)
    monkeypatch.undo()
    monkeypatch.setenv("LOCALAPPDATA", str(tmp_path / "localappdata"))
    messages, errors = [], []
    monkeypatch.setattr(tk_root, "refresh_pending", lambda: None, raising=False)
    monkeypatch.setattr(gui.messagebox, "showinfo", lambda *a, **k: messages.append(a))
    monkeypatch.setattr(gui.messagebox, "showerror", lambda *a, **k: errors.append(a))
    panel = RecoveryPanel(tk_root)
    panel.withdraw()
    manager = tasks_for(tk_root)
    try:
        panel.run(record)
        deadline = time.monotonic() + 20
        while manager.busy and time.monotonic() < deadline:
            tk_root.update()
            time.sleep(.01)
        assert not manager.busy
        assert messages and not errors
        assert recovery.pending() == []
        assert destination.exists()
    finally:
        if manager.thread:
            manager.thread.join(timeout=6)
        panel.destroy()
