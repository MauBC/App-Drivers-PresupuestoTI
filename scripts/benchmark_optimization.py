"""Benchmark reproducible; compara funciones contra el baseline 55747c5.

Ejecutar: python -m scripts.benchmark_optimization --repeats 3
No modifica Excel de usuario. Los tiempos no son assertions de tests.
"""
import argparse
import ast
import json
from statistics import median
import subprocess
from time import perf_counter

import pandas as pd
from openpyxl import Workbook

from src.core import search_executor as search
from src.core import driver_executor as driver


def baseline_function(module, path, name):
    source = subprocess.check_output(["git", "show", f"55747c5:{path}"], encoding="utf-8-sig")
    parsed = ast.parse(source)
    node = next(n for n in parsed.body if isinstance(n, ast.FunctionDef) and n.name == name)
    namespace = vars(module).copy()
    nodes = [node]
    if name == "write_dataframe_to_workbook_sheet":
        # El benchmark necesita el helper destructivo original solo en su referencia.
        helper = next(n for n in parsed.body if isinstance(n, ast.FunctionDef) and n.name == "delete_sheet_if_exists")
        nodes.insert(0, helper)
    exec(compile(ast.Module(body=nodes, type_ignores=[]), path, "exec"), namespace)
    return namespace[name]


def measure(function, repeats):
    times = []
    result = None
    for _ in range(repeats):
        started = perf_counter()
        result = function()
        times.append(perf_counter() - started)
    return median(times), result


def record(name, old, new):
    return {"scenario": name, "baseline_seconds": round(old, 4), "optimized_seconds": round(new, 4),
            "speedup": round(old / new, 2)}


def benchmark(repeats):
    previous_match = baseline_function(search, "src/core/search_executor.py", "run_similarity_search")
    previous_format = baseline_function(driver, "src/core/driver_executor.py", "write_dataframe_to_workbook_sheet")
    column = {"alias": "Nombre", "indice_0": 0, "indice_1": 1}
    config = {"template": {"columnas": [column], "columna_busqueda": column},
              "base": {"columnas": [column], "columna_busqueda": column,
                       "archivo": {"nombre": "maestro.xlsx"}, "hoja": {"nombre_detectado": "Datos"}},
              "match": {"tipo_busqueda": "similaridad", "score_exacto": 95, "score_aproximado": 70},
              "salida": {"incluir_columnas_tecnicas": True}}
    base = [{"Nombre": f"PERSONA NOMBRE {i:05d}"} for i in range(6000)]
    results = []
    for unique in (700, 50):
        templates = [{"Nombre": f"PERS0NA NOMBRE {i % unique:05d}"} for i in range(700)]
        before, old_rows = measure(lambda: previous_match(config, templates, base), repeats)
        after, new_rows = measure(lambda: search.run_similarity_search(config, templates, base), repeats)
        assert new_rows == old_rows, "Cambio en resultado de similaridad"
        results.append(record(f"similaridad_700_vs_6000_{unique}_consultas_distintas", before, after))

    df = pd.DataFrame({
        "FILA_ORIGEN": range(2, 6002), "DNI": [f"{i:08d}" for i in range(6000)],
        "CECO": ["A"] * 6000, "CECO2": ["B"] * 6000,
        "PORCENTAJE": [.25] * 6000, "PRECIO": [100.] * 6000,
        "PRECIO_DISTRIBUIDO": [25.] * 6000, "OBSERVACION": ["DNI encontrado"] * 6000,
    })
    def format_sheet(function):
        wb = Workbook()
        try:
            function(wb, "DRIVER", df)
            return wb
        except Exception:
            wb.close()
            raise
    def measured_format(function):
        times = []
        last = None
        for _ in range(repeats):
            if last is not None:
                last.close()
            started = perf_counter()
            last = format_sheet(function)
            times.append(perf_counter() - started)
        return median(times), last
    before, original = measured_format(previous_format)
    after, optimized = measured_format(driver.write_dataframe_to_workbook_sheet)
    try:
        a, b = original["DRIVER"], optimized["DRIVER"]
        assert list(a.values) == list(b.values)
        assert a.freeze_panes == b.freeze_panes and a.auto_filter.ref == b.auto_filter.ref
        assert {k: v.width for k, v in a.column_dimensions.items()} == {k: v.width for k, v in b.column_dimensions.items()}
        for old_row, new_row in zip(a.iter_rows(), b.iter_rows()):
            assert [(c.number_format, c._style) for c in old_row] == [(c.number_format, c._style) for c in new_row]
    finally:
        original.close()
        optimized.close()
    results.append(record("construccion_hoja_driver_6000_filas_8_columnas", before, after))
    return {"baseline_commit": "55747c5", "repeats": repeats, "statistic": "median", "results_identical": True,
            "measurements": results}


if __name__ == "__main__":
    parser = argparse.ArgumentParser()
    parser.add_argument("--repeats", type=int, default=3)
    args = parser.parse_args()
    if args.repeats < 1:
        parser.error("repeats debe ser positivo")
    print(json.dumps(benchmark(args.repeats), indent=2))
